# -*- coding: utf-8 -*-
r"""
Module TỪ VỰNG PREMIUM - cắm vào fix.py.
Tự sinh bộ thủ + mẹo nhớ từ module vocab_data/.
"""

import os
import re
import json

import openpyxl

from vocab_data.mnemonic_generator import generate_mnemonic
from vocab_data.radical_analyzer import get_radical_for_word

_SIMILAR_CHARS = {}
_SIMILAR_CHARS_LOADED = False
def _normalize_hsk_with_sheet(hsk, sheet_name=""):
    hsk_str = str(hsk or "").strip().upper()
    sheet_str = str(sheet_name or "").strip()

    hsk_clean = hsk_str
    m = re.match(r'^(HSK)\s*(\d+.*)$', hsk_str, re.IGNORECASE)
    if m:
        hsk_clean = m.group(1).upper() + " " + m.group(2).strip()

    if "7-9" in hsk_clean or "7 - 9" in hsk_clean:
        m_sheet = re.search(r'\(\s*(\d+)\s*\)', sheet_str)
        if m_sheet:
            sheet_num = m_sheet.group(1)
            if sheet_num == "2" and re.search(r'\(\s+2\s*\)', sheet_str):
                hsk_clean = "HSK 7-9 ( 2)"
            else:
                hsk_clean = "HSK 7-9 (" + sheet_num + ")"

    return hsk_clean


def _find_similar_chars_files():
    base = os.path.dirname(os.path.abspath(__file__))
    candidates = [
        os.path.join(base, "vocab_data"),
        os.path.join(os.getcwd(), "vocab_data"),
        "vocab_data",
    ]
    files = []
    for folder in candidates:
        if os.path.isdir(folder):
            for fname in os.listdir(folder):
                if fname.startswith("similar_chars") and fname.endswith(".json"):
                    files.append(os.path.join(folder, fname))
            if files:
                break
    return files


def load_similar_chars():
    global _SIMILAR_CHARS, _SIMILAR_CHARS_LOADED
    if _SIMILAR_CHARS_LOADED:
        return _SIMILAR_CHARS
    _SIMILAR_CHARS_LOADED = True

    files = _find_similar_chars_files()
    if not files:
        print("[VOCAB] [INFO] Khong co similar_chars*.json")
        return {}

    for path in files:
        try:
            with open(path, "r", encoding="utf-8") as f:
                data = json.load(f)
            if isinstance(data, dict):
                _SIMILAR_CHARS.update(data)
                print("[VOCAB] [OK] Load similar_chars: "
                      + os.path.basename(path)
                      + " (" + str(len(data)) + " entries)")
        except Exception as e:
            print("[VOCAB] [WARN] Loi doc " + path + ": " + str(e))

    print("[VOCAB] [OK] Tong similar_chars: "
          + str(len(_SIMILAR_CHARS)) + " chu")
    return _SIMILAR_CHARS


def _get_similar_chars(zh):
    if not zh or not _SIMILAR_CHARS:
        return []
    return _SIMILAR_CHARS.get(zh, [])


_AI_MNEMONICS = {}
_AI_MNEMONICS_LOADED = False


def _find_ai_mnemonics_file():
    candidates = [
        os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "ai_mnemonics.json"),
        os.path.join(os.getcwd(), "data", "ai_mnemonics.json"),
        "data/ai_mnemonics.json",
    ]
    for p in candidates:
        if os.path.exists(p):
            return p
    return None


def load_ai_mnemonics():
    global _AI_MNEMONICS, _AI_MNEMONICS_LOADED
    if _AI_MNEMONICS_LOADED:
        return _AI_MNEMONICS
    _AI_MNEMONICS_LOADED = True

    path = _find_ai_mnemonics_file()
    if not path:
        print("[VOCAB] [INFO] Khong co ai_mnemonics.json - dung meo tinh")
        return {}

    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, dict):
            _AI_MNEMONICS = data
            print("[VOCAB] [OK] Da load " + str(len(data)) + " meo nho AI tu: " + path)
    except Exception as e:
        print("[VOCAB] [WARN] Loi doc ai_mnemonics.json: " + str(e))
    return _AI_MNEMONICS


def _get_ai_mnemonic(hsk, stt, zh):
    if not _AI_MNEMONICS:
        return ""

    hsk_str = str(hsk or "").strip()
    stt_str = str(stt or "").strip()
    zh_str = str(zh or "").strip()

    # ═══════════════════════════════════════════════════════════
    #  TẠO CÁC BIẾN THỂ CHỮ HÁN
    #  VD: "爸爸 | 爸" → ["爸爸 | 爸", "爸爸", "爸"]
    #      "爱好"     → ["爱好"]
    # ═══════════════════════════════════════════════════════════
    zh_variants = [zh_str]

    # Nếu có dấu " | " (space + gạch + space)
    if " | " in zh_str:
        parts = zh_str.split(" | ")
        for p in parts:
            p = p.strip()
            if p and p not in zh_variants:
                zh_variants.append(p)

    # Nếu có dấu "|" (không space)
    elif "|" in zh_str:
        parts = zh_str.split("|")
        for p in parts:
            p = p.strip()
            if p and p not in zh_variants:
                zh_variants.append(p)

    # Nếu có space giữa các chữ (VD: "爸 爸")
    if " " in zh_str and " | " not in zh_str:
        no_space = zh_str.replace(" ", "")
        if no_space and no_space not in zh_variants:
            zh_variants.append(no_space)

    # ═══════════════════════════════════════════════════════════
    #  TẠO CÁC BIẾN THỂ HSK
    #  VD: "HSK 1" → ["HSK 1", "HSK1", "hsk 1", "hsk1"]
    # ═══════════════════════════════════════════════════════════
    m = re.match(r'^(HSK)\s*(\d+.*)$', hsk_str, re.IGNORECASE)
    hsk_spaced = (m.group(1).upper() + " " + m.group(2)) if m else hsk_str
    hsk_no_space = hsk_str.replace(" ", "")
    hsk_spaced_no_space = hsk_spaced.replace(" ", "")

    hsk_variants = [
        hsk_str,
        hsk_spaced,
        hsk_no_space,
        hsk_spaced_no_space,
        hsk_str.lower(),
        hsk_spaced.lower(),
        hsk_no_space.lower(),
    ]

    # ═══════════════════════════════════════════════════════════
    #  TẠO CANDIDATES (HSK × ZH)
    # ═══════════════════════════════════════════════════════════
    candidates = []
    for hsk_var in hsk_variants:
        for zh_var in zh_variants:
            candidates.append(hsk_var + "|" + stt_str + "|" + zh_var)

    # ═══════════════════════════════════════════════════════════
    #  THỬ MATCH TỪNG CANDIDATE
    # ═══════════════════════════════════════════════════════════
    for key in candidates:
        if key in _AI_MNEMONICS:
            return _AI_MNEMONICS[key]

    # ═══════════════════════════════════════════════════════════
    #  FALLBACK: DUYỆT TOÀN BỘ DICT
    # ═══════════════════════════════════════════════════════════
    target_hsk = hsk_no_space.lower()
    target_stt = stt_str
    target_zh_set = set(v.strip() for v in zh_variants)

    for k, v in _AI_MNEMONICS.items():
        parts = k.split("|")
        if len(parts) != 3:
            continue

        key_hsk = parts[0].replace(" ", "").lower()
        key_stt = parts[1].strip()
        key_zh = parts[2].strip()

        if key_hsk != target_hsk:
            continue
        if key_stt != target_stt:
            continue
        if key_zh in target_zh_set:
            return v

    return ""


try:
    import jieba
    HAS_JIEBA = True
    print("[VOCAB] OK - Da load jieba")
except ImportError:
    HAS_JIEBA = False
    print("[VOCAB] WARN - Khong co jieba, se tach tung chu don")


def _split_chinese_words(text):
    if not text:
        return []
    text = text.strip()
    if not text:
        return []

    words = []
    if HAS_JIEBA:
        try:
            for seg in jieba.cut(text):
                seg = seg.strip()
                if seg and re.match(r'^[\u4e00-\u9fff]+$', seg):
                    words.append(seg)
            if words:
                return words
        except Exception as e:
            print("      [WARN] jieba error: " + str(e))

    for c in text:
        if '\u4e00' <= c <= '\u9fff':
            words.append(c)

    return words


def _clean(s):
    if s is None:
        return ""
    return (str(s)
            .replace('\n', ' ')
            .replace('\r', ' ')
            .replace('\t', ' ')
            .strip())


def _clean_pinyin(s):
    if s is None:
        return ""
    text = _clean(s).strip('/').strip()
    return re.sub(r'\s+', ' ', text)


def _normalize_hsk(sheet_name):
    if not sheet_name:
        return ""
    s = sheet_name.strip()
    s = re.sub(r'\s+', '', s)
    s = re.sub(r'\(\d+\)$', '', s)
    return s


def _is_hsk_sheet(name):
    return bool(name and name.strip().upper().startswith("HSK"))


def _parse_radical_raw(raw):
    if not raw:
        return None
    text = _clean(raw)
    if not text:
        return None
    if '|' in text:
        parts = [p.strip() for p in text.split('|')]
        if len(parts) >= 2:
            return {
                "zh": parts[0],
                "pinyin": parts[1] if len(parts) > 1 else "",
                "strokes": parts[2] if len(parts) > 2 else "",
                "meaning": parts[3] if len(parts) > 3 else "",
            }
    return {"zh": "", "pinyin": "", "strokes": "", "meaning": text}


def _parse_vi_du_from_mnemonic(mnemonic):
    if not mnemonic:
        return []

    m = re.search(
        r'📎\s*Ví dụ:?\s*(.*?)(?=\n\s*🔗|\n\s*💡|\n\s*📌|\n\s*🎬|\Z)',
        mnemonic,
        re.DOTALL
    )
    if not m:
        return []

    block = m.group(1).strip()
    if not block:
        return []

    result = []

    parts = re.split(r'\s*[/;]\s*', block)

    for part in parts:
        part = part.strip()
        if not part:
            continue

        m2 = re.match(
            r'^([\u4e00-\u9fff，。！？、；：""''（）]+)\s*'
            r'(?:[\(\（]([^\)\）]+)[\)\）])?'
            r'(?:\s*[-—–]\s*(.+))?$',
            part
        )

        if m2:
            zh = m2.group(1).strip()
            py = (m2.group(2) or "").strip()
            vi = (m2.group(3) or "").strip()

            if zh:
                result.append({
                    "zh": zh,
                    "pinyin": py,
                    "vi": vi
                })

    return result


def _detect_data_start_row(ws, fallback_row=3, col_zh=1):
    HEADER_KEYWORDS = (
        'stt', '汉字', '汉语', 'pinyin', 'từ', 'hsk', '发音',
        'nghĩa', 'loại từ', 'tt', 'số tt', 'chữ hán', 'tiếng trung'
    )

    try:
        for r_idx, row in enumerate(
            ws.iter_rows(min_row=1, max_row=6, values_only=True),
            start=1
        ):
            if not row:
                continue

            cell_val = ""
            if len(row) > col_zh and row[col_zh] is not None:
                cell_val = str(row[col_zh]).strip().lower()

            if not cell_val:
                continue

            is_header = False
            for kw in HEADER_KEYWORDS:
                if kw in cell_val:
                    is_header = True
                    break
            if is_header:
                continue

            has_hanzi = any('\u4e00' <= c <= '\u9fff' for c in cell_val)
            if has_hanzi:
                print("      [AUTO] Data start row: " + str(r_idx))
                return r_idx

        print("      [AUTO] Fallback start row: " + str(fallback_row))
        return fallback_row

    except Exception as e:
        print("      [WARN] _detect_data_start_row error: " + str(e))
        return fallback_row




def read_vocab_excel(excel_file, start_row=3):
    print("\n[VOCAB] Dang doc: " + excel_file)
    if not os.path.exists(excel_file):
        print("   [X] Khong tim thay file")
        return []

    load_ai_mnemonics()
    load_similar_chars()

    try:
        wb = openpyxl.load_workbook(excel_file, data_only=True, read_only=True)
    except Exception as e:
        print("   [X] Loi mo file: " + str(e))
        return []

    target_sheets = [s for s in wb.sheetnames if _is_hsk_sheet(s)]
    if not target_sheets:
        print("   [!] Khong co sheet HSK nao")
        wb.close()
        return []

    print("   [OK] " + str(len(target_sheets)) + " sheet: " + ", ".join(target_sheets))

    all_data = []
    total_mnemonic_ai = 0
    total_mnemonic_static = 0
    total_radical_generated = 0

    for sheet_name in target_sheets:
        try:
            ws = wb[sheet_name]
            hsk = _normalize_hsk(sheet_name)
            actual_start_row = _detect_data_start_row(ws, start_row)

            rows, n_ai, n_static, n_rd = _read_vocab_sheet(
                ws, hsk, sheet_name, actual_start_row
            )
            all_data.extend(rows)
            total_mnemonic_ai += n_ai
            total_mnemonic_static += n_static
            total_radical_generated += n_rd

            msg = "      " + sheet_name + ": " + str(len(rows)) + " tu"
            if n_ai > 0:
                msg += " (" + str(n_ai) + " meo AI)"
            if n_static > 0:
                msg += " (" + str(n_static) + " meo tinh)"
            if n_rd > 0:
                msg += " (" + str(n_rd) + " bo thu)"
            print(msg)
        except Exception as e:
            print("      [X] Loi sheet " + sheet_name + ": " + str(e))

    wb.close()
    print("   [OK] Tong: " + str(len(all_data)) + " tu vung")
    if total_mnemonic_ai > 0:
        print("   Meo nho AI: " + str(total_mnemonic_ai) + " tu")
    if total_mnemonic_static > 0:
        print("   Meo nho tinh (fallback): " + str(total_mnemonic_static) + " tu")
    if total_radical_generated > 0:
        print("   Tu tim bo thu: " + str(total_radical_generated) + " tu")

    return all_data


def _read_vocab_sheet(ws, hsk, sheet_name, start_row):
    is_hsk79 = ('7' in hsk and '9' in hsk) or ('HSK7' in hsk) or ('HSK8' in hsk) or ('HSK9' in hsk)

    if is_hsk79:
        COL_STT = 0
        COL_ZH = 1
        COL_PINYIN = 2
        COL_LOAI_TU = -1
        COL_VI = 4
        COL_VI_DU_ZH = -1
        COL_VI_DU_PINYIN = -1
        COL_VI_DU_VI = -1
        COL_MNEMONIC = -1
        COL_RADICAL = -1
        print("      [COLS] HSK 7-9 mode (nghia o cot E)")
    else:
        COL_STT = 0
        COL_ZH = 1
        COL_PINYIN = 2
        COL_LOAI_TU = 4
        COL_VI = 5
        COL_VI_DU_ZH = 8
        COL_VI_DU_PINYIN = 9
        COL_VI_DU_VI = 10
        COL_MNEMONIC = 11
        COL_RADICAL = 12
        print("      [COLS] HSK 1-6 mode (nghia o cot F)")

    data = []
    empty_count = 0
    n_mnemonic_ai = 0
    n_mnemonic_static = 0
    n_radical_generated = 0

    sheet_clean = re.sub(r'[^A-Za-z0-9]+', '-', sheet_name).strip('-')

    for row in ws.iter_rows(min_row=start_row, values_only=True):
        if not row:
            empty_count += 1
            if empty_count > 30:
                break
            continue
        if len(row) <= COL_ZH:
            continue

        zh = _clean(row[COL_ZH]) if COL_ZH < len(row) else ""
        if not zh:
            empty_count += 1
            if empty_count > 30:
                break
            continue
        empty_count = 0

        if zh.lower() in ("từ tiếng trung", "汉字", "từ vựng", "từ", "hsk",
                          "汉语", "chữ hán", "tiếng trung"):
            continue

        stt_raw_val = row[COL_STT] if COL_STT < len(row) and row[COL_STT] is not None else ""
        stt_str = str(stt_raw_val).strip()
        stt_unique = (sheet_clean + "-" + stt_str) if stt_str else ""

        pinyin = _clean_pinyin(row[COL_PINYIN]) if COL_PINYIN >= 0 and COL_PINYIN < len(row) else ""
        vi = _clean(row[COL_VI]) if COL_VI >= 0 and COL_VI < len(row) else ""

        hsk_normalized = _normalize_hsk_with_sheet(hsk, sheet_name)

        mnemonic = ""
        source = "none"

        if COL_MNEMONIC >= 0 and COL_MNEMONIC < len(row):
            mnemonic = _clean(row[COL_MNEMONIC])
            if mnemonic:
                source = "excel"

        if not mnemonic:
            ai_text = _get_ai_mnemonic(hsk_normalized, stt_str, zh)
            if ai_text:
                mnemonic = ai_text
                source = "ai"
                n_mnemonic_ai += 1

        if not mnemonic:
            try:
                mnemonic = generate_mnemonic(zh, pinyin, vi)
                if mnemonic:
                    source = "static"
                    n_mnemonic_static += 1
            except Exception:
                mnemonic = ""

        vi_du_list = _parse_vi_du_from_mnemonic(mnemonic)

        radical = None
        if COL_RADICAL >= 0 and COL_RADICAL < len(row):
            radical = _parse_radical_raw(row[COL_RADICAL])
        if not radical:
            try:
                radical = get_radical_for_word(zh)
                if radical:
                    n_radical_generated += 1
            except Exception:
                radical = None

        vi_du_zh = _clean(row[COL_VI_DU_ZH]) if COL_VI_DU_ZH >= 0 and COL_VI_DU_ZH < len(row) else ""
        vi_du_pinyin = _clean_pinyin(row[COL_VI_DU_PINYIN]) if COL_VI_DU_PINYIN >= 0 and COL_VI_DU_PINYIN < len(row) else ""
        vi_du_vi = _clean(row[COL_VI_DU_VI]) if COL_VI_DU_VI >= 0 and COL_VI_DU_VI < len(row) else ""
        vi_du_words = _split_chinese_words(vi_du_zh)

        data.append({
            "stt": stt_unique,
            "stt_original": stt_str,
            "hsk": hsk_normalized,
            "topic": "Từ vựng",
            "subject": _clean(row[COL_LOAI_TU]) if COL_LOAI_TU >= 0 and COL_LOAI_TU < len(row) else "",
            "vi": vi,
            "zh": zh,
            "pinyin": pinyin,
            "vi_du_zh": vi_du_zh,
            "vi_du_pinyin": vi_du_pinyin,
            "vi_du_vi": vi_du_vi,
            "vi_du_words": vi_du_words,
            "mnemonic": mnemonic,
            "mnemonic_source": source,
            "vi_du_list": vi_du_list,
            "radical": radical,
            "source_sheet": sheet_name,
        })

    return data, n_mnemonic_ai, n_mnemonic_static, n_radical_generated
def build_vocab_css(vocab_id="tu-vung"):
    css = r"""
.ds-btn[data-dataset="__VOCAB_ID__"] {
    background: linear-gradient(135deg,
        rgba(251, 191, 36, .12) 0%,
        rgba(245, 158, 11, .08) 50%,
        rgba(8, 145, 178, .06) 100%);
    border: 2px solid rgba(245, 158, 11, .4);
    color: #92400e;
    position: relative;
    overflow: visible;
    font-weight: 800;
    padding-left: 2.5rem;
    box-shadow: 0 0 0 1px rgba(245, 158, 11, .2), 0 2px 8px rgba(245, 158, 11, .15);
}
.ds-btn[data-dataset="__VOCAB_ID__"]::before {
    content: '👑';
    position: absolute;
    left: .6rem;
    top: 50%;
    transform: translateY(-50%);
    font-size: 1.1rem;
    line-height: 1;
    z-index: 3;
    filter: drop-shadow(0 2px 4px rgba(245, 158, 11, .8));
    animation: crownTabFloat 3s ease-in-out infinite;
    pointer-events: none;
}
@keyframes crownTabFloat {
    0%, 100% { transform: translateY(-50%) rotate(0deg); }
    50%      { transform: translateY(-50%) rotate(-10deg) scale(1.1); }
}
.ds-btn[data-dataset="__VOCAB_ID__"] > i:first-child { display: none; }
.ds-btn[data-dataset="__VOCAB_ID__"]:hover {
    background: linear-gradient(135deg,
        rgba(251, 191, 36, .22) 0%,
        rgba(245, 158, 11, .15) 50%,
        rgba(8, 145, 178, .1) 100%);
    border-color: #f59e0b;
    color: #78350f;
    box-shadow: 0 0 0 2px rgba(245, 158, 11, .6), 0 4px 16px rgba(245, 158, 11, .35);
    transform: translateY(-1px);
}
.ds-btn[data-dataset="__VOCAB_ID__"].active {
    background: linear-gradient(135deg, #fbbf24 0%, #f59e0b 30%, #0891b2 100%);
    color: #fff;
    border-color: transparent;
    box-shadow: 0 0 0 2px rgba(245, 158, 11, .7), 0 6px 22px rgba(245, 158, 11, .5);
    text-shadow: 0 1px 2px rgba(0, 0, 0, .2);
}
.ds-btn[data-dataset="__VOCAB_ID__"].active::before {
    filter: drop-shadow(0 0 8px rgba(255, 255, 255, .9));
}
[data-theme="dark"] .ds-btn[data-dataset="__VOCAB_ID__"] {
    background: linear-gradient(135deg,
        rgba(251, 191, 36, .2) 0%,
        rgba(245, 158, 11, .14) 50%,
        rgba(8, 145, 178, .1) 100%);
    color: #fcd34d;
    border-color: rgba(245, 158, 11, .5);
}
[data-theme="dark"] .ds-btn[data-dataset="__VOCAB_ID__"].active {
    background: linear-gradient(135deg, #d97706, #b45309 30%, #0e7490);
    color: #fff;
}
.ds-btn[data-dataset="__VOCAB_ID__"] .ds-vocab-badge {
    position: absolute;
    top: -10px; right: -8px;
    padding: .2rem .5rem;
    border-radius: 50px;
    background: linear-gradient(135deg, #6366f1 0%, #7c3aed 50%, #a855f7 100%);
    color: #fff;
    font-size: .55rem;
    font-weight: 900;
    letter-spacing: .5px;
    text-transform: uppercase;
    box-shadow: 0 2px 8px rgba(124, 58, 237, .6), 0 0 0 2px var(--surface);
    animation: premiumBadgePulse 2s ease-in-out infinite;
    z-index: 10;
    pointer-events: none;
    line-height: 1.2;
    white-space: nowrap;
}
.ds-btn[data-dataset="__VOCAB_ID__"] .ds-vocab-badge::before { content: '💎 '; }
@keyframes premiumBadgePulse {
    0%, 100% { transform: scale(1); }
    50%      { transform: scale(1.1); }
}
.ds-btn[data-dataset="__VOCAB_ID__"].active .ds-vocab-badge {
    background: linear-gradient(135deg, #fbbf24, #f59e0b);
    color: #1e1b4b;
    animation: none;
}
.ds-btn[data-dataset="__VOCAB_ID__"].vocab-locked {
    background: linear-gradient(135deg, rgba(220, 38, 38, .08) 0%, rgba(251, 191, 36, .06) 100%);
    border-color: rgba(220, 38, 38, .4);
    color: #991b1b;
    opacity: .85;
}
.ds-btn[data-dataset="__VOCAB_ID__"].vocab-locked::before {
    filter: drop-shadow(0 2px 4px rgba(220, 38, 38, .6)) grayscale(.5);
    opacity: .7;
}
.ds-btn[data-dataset="__VOCAB_ID__"].vocab-locked:hover { opacity: 1; border-color: #dc2626; }
.ds-btn[data-dataset="__VOCAB_ID__"].vocab-locked .ds-vocab-badge { display: none; }
.ds-btn[data-dataset="__VOCAB_ID__"] .vocab-lock-icon {
    position: absolute;
    top: -8px; right: -6px;
    width: 22px; height: 22px;
    border-radius: 50%;
    background: linear-gradient(135deg, #dc2626, #b91c1c);
    color: #fff;
    display: flex; align-items: center; justify-content: center;
    font-size: .62rem;
    box-shadow: 0 2px 8px rgba(220, 38, 38, .6), 0 0 0 2px var(--surface);
    z-index: 11;
    animation: lockPulse 2.5s ease-in-out infinite;
}
@keyframes lockPulse {
    0%, 100% { transform: scale(1); }
    50%      { transform: scale(1.12); }
}

.card-radical {
    margin-top: .7rem;
    padding: .7rem .85rem;
    background: var(--surface-2);
    border: 1px solid var(--border);
    border-radius: 10px;
    display: flex; flex-direction: column; gap: .55rem;
    animation: blockIn .35s ease-out;
}
@keyframes blockIn {
    from { opacity: 0; transform: translateY(-3px); }
    to   { opacity: 1; transform: translateY(0); }
}
[data-theme="dark"] .card-radical {
    background: rgba(139, 92, 246, .08);
    border-color: rgba(139, 92, 246, .2);
}
.card-radical-label {
    display: flex; align-items: center; gap: .4rem;
    font-size: .72rem; font-weight: 800;
    color: #6d28d9; letter-spacing: .5px;
    text-transform: uppercase;
}
[data-theme="dark"] .card-radical-label { color: #c4b5fd; }
.card-radical-body { display: flex; align-items: center; gap: .75rem; }
.card-radical-box {
    width: 52px; height: 52px;
    border: 1.5px solid var(--border-strong);
    border-radius: 6px;
    display: flex; align-items: center; justify-content: center;
    font-family: var(--font-zh);
    font-size: 1.65rem; font-weight: 500;
    color: var(--text);
    flex-shrink: 0;
    background: var(--surface);
    position: relative;
}
.card-radical-box::before {
    content: '';
    position: absolute; inset: 0;
    background-image:
        linear-gradient(to right, transparent 49%, var(--border) 49%, var(--border) 51%, transparent 51%),
        linear-gradient(to bottom, transparent 49%, var(--border) 49%, var(--border) 51%, transparent 51%);
    opacity: .35; pointer-events: none; border-radius: 6px;
}
.card-radical-info {
    flex: 1; min-width: 0;
    display: flex; flex-direction: column; gap: .2rem;
}
.card-radical-name {
    font-size: .88rem; font-weight: 700;
    color: var(--text); line-height: 1.3;
}
.card-radical-name .pinyin {
    font-weight: 500; font-style: italic;
    color: #8b5cf6;
}
[data-theme="dark"] .card-radical-name .pinyin { color: #c4b5fd; }
.card-radical-meaning {
    font-size: .8rem; color: var(--text-2); line-height: 1.4;
}

.card-mnemonic {
    margin-top: .7rem;
    padding: .7rem .85rem;
    background: var(--surface-2);
    border: 1px solid var(--border);
    border-radius: 10px;
    display: flex; flex-direction: column; gap: .5rem;
    animation: blockIn .35s ease-out;
}
[data-theme="dark"] .card-mnemonic {
    background: rgba(245, 158, 11, .08);
    border-color: rgba(245, 158, 11, .2);
}
.card-mnemonic-label {
    display: flex; align-items: center; gap: .4rem;
    font-size: .72rem; font-weight: 800;
    color: #b45309; letter-spacing: .5px;
    text-transform: uppercase;
}
[data-theme="dark"] .card-mnemonic-label { color: #fcd34d; }
.card-mnemonic-body {
    font-size: .84rem; color: var(--text-2);
    line-height: 1.55; white-space: pre-line;
}
.card-mnemonic-body .char-zh {
    font-family: var(--font-zh);
    font-size: 1rem; font-weight: 600;
    color: var(--text); padding: 0 .15rem;
}
.card-mnemonic-body .arrow {
    color: #f59e0b; font-weight: 800; padding: 0 .2rem;
}
.card-mnemonic-body .hint {
    color: #b45309; font-weight: 700;
}
[data-theme="dark"] .card-mnemonic-body .hint { color: #fcd34d; }
/* ========== CỤM HÁN BẤM ĐƯỢC TRONG MẸO NHỚ ========== */
/* Style y hệt .card-example-zh nhưng dùng cho câu KHÔNG nằm trong block 📎 Ví dụ */
.card-mnemonic-body .card-char-speakable {
    font-family: var(--font-zh);
    font-weight: 600;
    color: var(--text);
    font-size: 1rem;
    cursor: pointer;
    user-select: none;
    -webkit-tap-highlight-color: transparent;
    border-bottom: 1.5px dotted rgba(8, 145, 178, .5);
    padding-bottom: 1px;
    transition: all .18s ease;
}
.card-mnemonic-body .card-char-speakable:hover {
    color: #0891b2;
    border-bottom-color: #0891b2;
    background: rgba(8, 145, 178, .1);
    border-radius: 4px;
    padding: 0 .15rem 1px;
}
.card-mnemonic-body .card-char-speakable:active {
    transform: scale(.96);
}
.card-mnemonic-body .card-char-speakable.speaking {
    color: #f59e0b;
    border-bottom-color: #f59e0b;
    animation: cardExampleSpeak 1s ease-in-out infinite;
}
[data-theme="dark"] .card-mnemonic-body .card-char-speakable {
    color: #e0f2fe;
    border-bottom-color: rgba(34, 211, 238, .5);
}
[data-theme="dark"] .card-mnemonic-body .card-char-speakable:hover {
    color: #22d3ee;
    border-bottom-color: #22d3ee;
    background: rgba(34, 211, 238, .12);
}
[data-theme="dark"] .card-mnemonic-body .card-char-speakable.speaking {
    color: #fbbf24;
    border-bottom-color: #fbbf24;
}

/* Fix cho vocab-mode compact (câu trong mẹo nhớ nhỏ hơn) */
body[data-vocab-mode="1"] .card-mnemonic-body .card-char-speakable {
    font-size: .9rem;
}
@media (max-width: 500px) {
    body[data-vocab-mode="1"] .card-mnemonic-body .card-char-speakable {
        font-size: .85rem;
    }
}
.card-mnemonic-body .card-example-inline {
    display: inline-flex;
    align-items: center;
    flex-wrap: wrap;
    gap: .35rem;
    margin-left: .3rem;
    padding: .15rem .4rem;
    background: rgba(8, 145, 178, .08);
    border-radius: 8px;
    border: 1px solid rgba(8, 145, 178, .15);
    vertical-align: middle;
}
.card-mnemonic-body .card-example-zh {
    font-family: var(--font-zh);
    font-weight: 700;
    color: var(--text);
    font-size: .95rem;
    cursor: pointer;
    user-select: none;
    -webkit-tap-highlight-color: transparent;
    border-bottom: 1.5px dotted rgba(8, 145, 178, .5);
    padding-bottom: 1px;
    transition: all .18s ease;
}
.card-mnemonic-body .card-example-zh:hover {
    color: #0891b2;
    border-bottom-color: #0891b2;
    background: rgba(8, 145, 178, .1);
    border-radius: 4px;
    padding: 0 .15rem 1px;
}
.card-mnemonic-body .card-example-zh:active {
    transform: scale(.96);
}
.card-mnemonic-body .card-example-zh.speaking {
    color: #f59e0b;
    border-bottom-color: #f59e0b;
    animation: cardExampleSpeak 1s ease-in-out infinite;
}
@keyframes cardExampleSpeak {
    0%, 100% { opacity: 1; }
    50%      { opacity: .55; }
}
.card-mnemonic-body .card-example-pinyin {
    font-style: italic;
    color: #0891b2;
    font-size: .78rem;
    font-weight: 500;
}
[data-theme="dark"] .card-mnemonic-body .card-example-pinyin {
    color: #22d3ee;
}
.card-mnemonic-body .card-example-vi {
    color: var(--text-2);
    font-size: .78rem;
}
[data-theme="dark"] .card-mnemonic-body .card-example-zh {
    color: #e0f2fe;
    border-bottom-color: rgba(34, 211, 238, .5);
}
[data-theme="dark"] .card-mnemonic-body .card-example-zh:hover {
    color: #22d3ee;
    border-bottom-color: #22d3ee;
    background: rgba(34, 211, 238, .12);
}
[data-theme="dark"] .card-mnemonic-body .card-example-zh.speaking {
    color: #fbbf24;
    border-bottom-color: #fbbf24;
}
[data-theme="dark"] .card-mnemonic-body .card-example-inline {
    background: rgba(34, 211, 238, .12);
    border-color: rgba(34, 211, 238, .25);
}
body[data-vocab-mode="1"] .card-mnemonic-body .card-example-inline {
    padding: .1rem .3rem;
    gap: .2rem;
}
body[data-vocab-mode="1"] .card-mnemonic-body .card-example-zh {
    font-size: .85rem;
}
body[data-vocab-mode="1"] .card-mnemonic-body .card-example-pinyin,
body[data-vocab-mode="1"] .card-mnemonic-body .card-example-vi {
    font-size: .7rem;
}

.pinyin-hl {
    background: linear-gradient(180deg, transparent 55%, rgba(250, 204, 21, .45) 55%);
    padding: 0 .08em; border-radius: 2px;
    color: #78350f; font-weight: 600;
}
[data-theme="dark"] .pinyin-hl {
    background: linear-gradient(180deg, transparent 55%, rgba(250, 204, 21, .3) 55%);
    color: #fde68a;
}

.pf-vocab-example {
    margin-top: 1rem;
    padding: .85rem 1rem;
    background: linear-gradient(135deg, rgba(8, 145, 178, .08), rgba(6, 182, 212, .04));
    border-left: 4px solid #0891b2;
    border-radius: 12px;
    display: flex;
    flex-direction: column;
    gap: .5rem;
    animation: pfVocabIn .4s ease-out;
}
@keyframes pfVocabIn {
    from { opacity: 0; transform: translateY(8px); }
    to   { opacity: 1; transform: translateY(0); }
}
[data-theme="dark"] .pf-vocab-example {
    background: linear-gradient(135deg, rgba(8, 145, 178, .18), rgba(6, 182, 212, .1));
    border-left-color: #22d3ee;
}
.pf-vocab-example-label {
    font-size: .68rem;
    font-weight: 800;
    color: #0891b2;
    text-transform: uppercase;
    letter-spacing: .5px;
    display: flex; align-items: center; gap: .4rem;
}
[data-theme="dark"] .pf-vocab-example-label { color: #22d3ee; }
.pf-vocab-example-zh {
    font-family: var(--font-zh);
    font-size: clamp(1.05rem, 2.2vw, 1.4rem);
    font-weight: 600;
    color: var(--text);
    line-height: 1.5;
    display: flex; align-items: center; gap: .5rem; flex-wrap: wrap;
}
.pf-vocab-example-zh .audio-btn-mini {
    width: 28px; height: 28px;
    border-radius: 50%; border: none;
    background: rgba(8, 145, 178, .15);
    color: #0891b2; cursor: pointer;
    display: inline-flex; align-items: center; justify-content: center;
    font-size: .75rem; transition: all .18s;
    flex-shrink: 0;
}
.pf-vocab-example-zh .audio-btn-mini:hover {
    background: #0891b2; color: #fff;
    transform: scale(1.15);
}
.pf-vocab-example-pinyin {
    font-size: clamp(.85rem, 1.4vw, 1rem);
    font-style: italic; color: #0891b2;
    font-weight: 500; line-height: 1.4;
}
[data-theme="dark"] .pf-vocab-example-pinyin { color: #22d3ee; }
.pf-vocab-example-vi {
    font-size: clamp(.9rem, 1.5vw, 1.05rem);
    color: var(--text-2);
    line-height: 1.5; font-weight: 500;
}

.vocab-upgrade-modal {
    position: fixed; inset: 0;
    background: rgba(15, 23, 42, .85);
    backdrop-filter: blur(6px);
    z-index: 6000; display: none;
    align-items: center; justify-content: center;
    padding: 1rem; animation: fadeIn .2s;
}
.vocab-upgrade-modal.show { display: flex; }
.vocab-upgrade-box {
    background: var(--surface);
    border-radius: 20px;
    max-width: 460px; width: 100%;
    max-height: calc(100vh - 2rem);
    overflow-y: auto;
    box-shadow: 0 20px 60px rgba(0,0,0,.4);
    animation: upgradeIn .35s cubic-bezier(.34, 1.56, .64, 1);
    text-align: center;
}
@keyframes upgradeIn {
    from { transform: scale(.9); opacity: 0; }
    to   { transform: scale(1); opacity: 1; }
}
.vocab-upgrade-header {
    padding: 1.75rem 1.5rem 1.25rem;
    background: linear-gradient(135deg, #6366f1 0%, #7c3aed 25%, #f59e0b 75%, #d97706 100%);
    border-radius: 20px 20px 0 0;
    color: #fff; position: relative; overflow: hidden;
}
.vocab-upgrade-icon {
    width: 70px; height: 70px;
    margin: 0 auto .85rem;
    border-radius: 50%;
    background: linear-gradient(135deg, #fbbf24, #f59e0b);
    display: flex; align-items: center; justify-content: center;
    font-size: 2rem;
    box-shadow: 0 8px 24px rgba(245, 158, 11, .5);
    animation: floatUp 3s ease-in-out infinite;
}
@keyframes floatUp {
    0%, 100% { transform: translateY(0); }
    50%      { transform: translateY(-6px); }
}
.vocab-upgrade-title {
    font-size: 1.35rem; font-weight: 900;
    margin-bottom: .35rem;
}
.vocab-upgrade-subtitle {
    font-size: .85rem; opacity: .95; line-height: 1.5;
}
.vocab-upgrade-price {
    display: inline-flex; align-items: center; gap: .4rem;
    margin-top: .85rem;
    padding: .55rem 1.15rem;
    background: linear-gradient(135deg, #fbbf24, #f59e0b);
    border: 2px solid rgba(255, 255, 255, .6);
    border-radius: 50px;
    font-weight: 900; font-size: 1.35rem; color: #fff;
    box-shadow: 0 6px 20px rgba(245, 158, 11, .6);
    animation: pricePulse 2.5s ease-in-out infinite;
}
@keyframes pricePulse {
    0%, 100% { transform: scale(1); }
    50%      { transform: scale(1.05); }
}
.vocab-upgrade-body {
    padding: 1.25rem 1.5rem;
    display: flex; flex-direction: column; gap: 1rem;
}
.vocab-upgrade-features {
    display: flex; flex-direction: column;
    gap: .55rem; text-align: left;
}
.vocab-upgrade-feature {
    display: flex; align-items: center; gap: .6rem;
    padding: .55rem .75rem;
    background: var(--surface-2);
    border-radius: 10px;
    font-size: .84rem; font-weight: 600;
    color: var(--text-2);
}
.vocab-upgrade-feature i {
    width: 24px; height: 24px;
    border-radius: 50%;
    background: rgba(8, 145, 178, .15);
    color: #0891b2;
    display: flex; align-items: center; justify-content: center;
    font-size: .75rem; flex-shrink: 0;
}
.vocab-upgrade-feature.highlight {
    background: linear-gradient(135deg, rgba(251, 191, 36, .15), rgba(245, 158, 11, .08));
    color: #92400e;
    border: 1.5px solid rgba(245, 158, 11, .4);
}
.vocab-upgrade-feature.highlight i {
    background: linear-gradient(135deg, #fbbf24, #f59e0b);
    color: #fff;
}
.vocab-upgrade-actions {
    display: flex; gap: .55rem; flex-wrap: wrap;
}
.vocab-upgrade-btn {
    flex: 1 1 160px;
    padding: .85rem 1rem;
    border-radius: 12px; border: none;
    font-size: .9rem; font-weight: 800;
    font-family: inherit; cursor: pointer;
    display: inline-flex; align-items: center; justify-content: center;
    gap: .45rem; transition: all .2s;
}
.vocab-upgrade-btn.primary {
    background: linear-gradient(135deg, #fbbf24, #f59e0b 50%, #ea580c);
    color: #fff;
    box-shadow: 0 6px 18px rgba(245, 158, 11, .5);
}
.vocab-upgrade-btn.primary:hover {
    transform: translateY(-2px);
    box-shadow: 0 10px 26px rgba(245, 158, 11, .7);
}
.vocab-upgrade-btn.secondary {
    background: var(--surface-2);
    color: var(--text-2);
    border: 1.5px solid var(--border);
}
.vocab-upgrade-btn.secondary:hover {
    border-color: var(--primary);
    color: var(--primary);
}

@media (max-width: 500px) {
    .ds-btn[data-dataset="__VOCAB_ID__"] { padding-left: 2.2rem; }
    .ds-btn[data-dataset="__VOCAB_ID__"]::before { font-size: 1rem; left: .5rem; }
    .card-radical, .card-mnemonic { padding: .6rem .7rem; margin-top: .6rem; }
    .card-radical-box { width: 46px; height: 46px; font-size: 1.45rem; }
    .card-mnemonic-body { font-size: .78rem; }
    .vocab-upgrade-icon { width: 60px; height: 60px; font-size: 1.65rem; }
    .vocab-upgrade-title { font-size: 1.15rem; }
    .vocab-upgrade-price { font-size: 1.15rem; }
    .pf-vocab-example { padding: .7rem .85rem; margin-top: .85rem; }
    .pf-vocab-example-zh { font-size: 1rem; }
    .pf-vocab-example-pinyin { font-size: .8rem; }
    .pf-vocab-example-vi { font-size: .85rem; }
}

.pf-chars-label {
    margin-top: .85rem;
    font-size: .7rem;
    font-weight: 700;
    color: #7c3aed;
    text-transform: uppercase;
    letter-spacing: .5px;
}
[data-theme="dark"] .pf-chars-label { color: #c4b5fd; }
.pf-chars-wrap {
    display: flex;
    flex-wrap: wrap;
    gap: .4rem;
    margin-top: .4rem;
}
.pf-char-btn {
    font-family: var(--font-zh);
    font-size: clamp(1.15rem, 2vw, 1.4rem);
    font-weight: 600;
    padding: .35rem .7rem;
    border-radius: 10px;
    border: 2px solid rgba(139, 92, 246, .3);
    background: linear-gradient(135deg, rgba(139, 92, 246, .08), rgba(124, 58, 237, .04));
    color: var(--text);
    cursor: pointer;
    transition: all .2s cubic-bezier(.34, 1.56, .64, 1);
    user-select: none;
    line-height: 1.2;
}
.pf-char-btn:hover {
    transform: translateY(-2px) scale(1.08);
    background: linear-gradient(135deg, rgba(139, 92, 246, .2), rgba(124, 58, 237, .1));
    border-color: #8b5cf6;
    box-shadow: 0 4px 12px rgba(139, 92, 246, .3);
}
.pf-char-btn:active {
    transform: translateY(0) scale(.98);
}
.pf-char-btn.active {
    background: linear-gradient(135deg, #8b5cf6, #7c3aed);
    color: #fff;
    border-color: #7c3aed;
    box-shadow: 0 6px 18px rgba(124, 58, 237, .5);
    transform: translateY(-2px) scale(1.1);
}
[data-theme="dark"] .pf-char-btn {
    border-color: rgba(167, 139, 250, .4);
    background: linear-gradient(135deg, rgba(139, 92, 246, .15), rgba(124, 58, 237, .08));
    color: #e9d5ff;
}
[data-theme="dark"] .pf-char-btn:hover {
    border-color: #a78bfa;
}

.pf-char-info {
    margin-top: .85rem;
    padding: .85rem 1rem;
    background: linear-gradient(135deg, rgba(139, 92, 246, .1), rgba(124, 58, 237, .05));
    border-left: 4px solid #8b5cf6;
    border-radius: 12px;
    display: flex;
    flex-direction: column;
    gap: .5rem;
    animation: pfCharInfoIn .3s ease-out;
}
@keyframes pfCharInfoIn {
    from { opacity: 0; transform: translateY(6px); }
    to   { opacity: 1; transform: translateY(0); }
}
[data-theme="dark"] .pf-char-info {
    background: linear-gradient(135deg, rgba(139, 92, 246, .2), rgba(124, 58, 237, .1));
    border-left-color: #a78bfa;
}
.pf-char-info-header {
    display: flex;
    align-items: baseline;
    gap: .6rem;
    padding-bottom: .5rem;
    border-bottom: 1px dashed rgba(139, 92, 246, .3);
    flex-wrap: wrap;
}
.pf-char-info-zh {
    font-family: var(--font-zh);
    font-size: 1.6rem;
    font-weight: 700;
    color: #7c3aed;
}
[data-theme="dark"] .pf-char-info-zh { color: #c4b5fd; }
.pf-char-info-pinyin {
    font-size: 1rem;
    font-style: italic;
    color: #8b5cf6;
    font-weight: 500;
}
[data-theme="dark"] .pf-char-info-pinyin { color: #a78bfa; }
.pf-char-info-line {
    display: flex;
    gap: .5rem;
    font-size: .85rem;
    line-height: 1.5;
    color: var(--text-2);
    align-items: flex-start;
    flex-wrap: wrap;
}
.pf-char-info-label {
    font-weight: 800;
    color: #7c3aed;
    white-space: nowrap;
    flex-shrink: 0;
}
[data-theme="dark"] .pf-char-info-label { color: #c4b5fd; }
.pf-char-info-mnemonic {
    white-space: pre-line;
    line-height: 1.6;
    color: var(--text);
    font-size: .82rem;
}
.pf-char-info-empty {
    padding: .75rem;
    text-align: center;
    color: var(--text-3);
    font-size: .85rem;
    font-style: italic;
}

.card-mnemonic-body .similar-hint-block {
    display: block;
    margin-top: .8rem;
    padding: .65rem .85rem;
    background: linear-gradient(135deg,
        rgba(245, 158, 11, .12) 0%,
        rgba(251, 191, 36, .08) 50%,
        rgba(239, 68, 68, .06) 100%);
    border-left: 3px solid #f59e0b;
    border-radius: 10px;
    font-size: .82rem;
    color: #92400e;
    line-height: 1.7;
    position: relative;
}
.card-mnemonic-body .similar-hint-block .similar-title {
    display: block;
    font-weight: 800;
    color: #b45309;
    margin-bottom: .5rem;
    font-size: .78rem;
    letter-spacing: .3px;
    text-transform: uppercase;
}
.card-mnemonic-body .similar-hint-block .similar-chars {
    display: flex;
    flex-wrap: wrap;
    align-items: center;
    gap: .35rem .5rem;
    margin-bottom: .35rem;
}
.card-mnemonic-body .similar-char-item {
    display: inline-flex;
    align-items: center;
    gap: .3rem;
    padding: .2rem .4rem .2rem .2rem;
    border-radius: 8px;
    background: rgba(220, 38, 38, .06);
    border: 1.5px solid rgba(220, 38, 38, .2);
    transition: all .18s ease;
    user-select: none;
}
.card-mnemonic-body .similar-char-item:hover {
    background: rgba(220, 38, 38, .1);
    border-color: rgba(220, 38, 38, .4);
}
.card-mnemonic-body .similar-char-btn-audio {
    width: 24px; height: 24px;
    border-radius: 50%;
    border: none;
    background: rgba(139, 92, 246, .15);
    color: #8b5cf6;
    cursor: pointer;
    display: inline-flex;
    align-items: center;
    justify-content: center;
    font-size: .65rem;
    transition: all .18s;
    flex-shrink: 0;
    padding: 0;
}
.card-mnemonic-body .similar-char-btn-audio:hover {
    background: #8b5cf6;
    color: #fff;
    transform: scale(1.15);
}
.card-mnemonic-body .similar-char-btn-audio.speaking {
    background: #dc2626;
    color: #fff;
    animation: audioPulse 1s infinite;
}
@keyframes audioPulse {
    0%, 100% { box-shadow: 0 0 0 0 rgba(220, 38, 38, .6); }
    50%      { box-shadow: 0 0 0 6px rgba(220, 38, 38, 0); }
}
.card-mnemonic-body .similar-char-main {
    display: inline-flex;
    align-items: baseline;
    gap: .25rem;
}
.card-mnemonic-body .similar-char-main .similar-char-zh {
    font-family: var(--font-zh);
    font-size: 1.05rem;
    font-weight: 700;
    color: #dc2626;
    line-height: 1;
}
.card-mnemonic-body .similar-char-main .similar-char-info-txt {
    font-size: .7rem;
    color: var(--text-2);
    font-weight: 500;
    line-height: 1;
    display: inline-flex;
    align-items: baseline;
    gap: .15rem;
}
.card-mnemonic-body .similar-char-main .similar-char-pinyin {
    font-style: italic;
    color: #8b5cf6;
    font-weight: 600;
    font-size: .72rem;
}
.card-mnemonic-body .similar-char-btn-jump {
    width: 24px; height: 24px;
    border-radius: 50%;
    border: none;
    background: rgba(220, 38, 38, .15);
    color: #dc2626;
    cursor: pointer;
    display: inline-flex;
    align-items: center;
    justify-content: center;
    font-size: .65rem;
    transition: all .18s;
    flex-shrink: 0;
    padding: 0;
}
.card-mnemonic-body .similar-char-btn-jump:hover {
    background: #dc2626;
    color: #fff;
    transform: translateX(3px);
}
.card-mnemonic-body .similar-sep {
    color: rgba(220, 38, 38, .5);
    font-weight: 700;
    font-size: .9rem;
    user-select: none;
}

.nav-back-btn {
    position: fixed;
    bottom: calc(16px + env(safe-area-inset-bottom));
    left: 16px;
    padding: .65rem 1rem;
    background: linear-gradient(135deg, #dc2626, #b91c1c);
    color: #fff;
    border: none;
    border-radius: 50px;
    font-family: inherit;
    font-size: .85rem;
    font-weight: 800;
    cursor: pointer;
    display: inline-flex;
    align-items: center;
    gap: .4rem;
    box-shadow: 0 8px 24px rgba(220, 38, 38, .5), 0 4px 12px rgba(0, 0, 0, .2);
    z-index: 1500;
    animation: navBackIn .35s cubic-bezier(.34, 1.56, .64, 1);
    transition: all .2s;
}
.nav-back-btn:hover {
    transform: translateY(-2px) scale(1.05);
    box-shadow: 0 12px 32px rgba(220, 38, 38, .7);
}
.nav-back-btn:active { transform: translateY(0) scale(.97); }
@keyframes navBackIn {
    from { opacity: 0; transform: translateY(20px) scale(.9); }
    to   { opacity: 1; transform: translateY(0) scale(1); }
}

.card.nav-highlight {
    animation: navCardPulse 2s ease-out;
    position: relative;
    z-index: 5;
}
.card.nav-highlight::before {
    content: '';
    position: absolute;
    inset: -3px;
    border-radius: 20px;
    border: 3px solid #dc2626;
    pointer-events: none;
    animation: navBorderPulse 2s ease-out;
}
@keyframes navCardPulse {
    0%, 100% { transform: translateY(0) scale(1); }
    15%      { transform: translateY(-4px) scale(1.02); }
    30%      { transform: translateY(0) scale(1); }
}
@keyframes navBorderPulse {
    0%   { opacity: 1; box-shadow: 0 0 0 0 rgba(220, 38, 38, .6); }
    50%  { opacity: .9; box-shadow: 0 0 0 12px rgba(220, 38, 38, 0); }
    100% { opacity: 0; box-shadow: 0 0 0 20px rgba(220, 38, 38, 0); }
}

[data-theme="dark"] .card-mnemonic-body .similar-hint-block {
    background: linear-gradient(135deg,
        rgba(245, 158, 11, .18) 0%,
        rgba(251, 191, 36, .12) 50%,
        rgba(239, 68, 68, .08) 100%);
    color: #fcd34d;
    border-left-color: #fbbf24;
}
[data-theme="dark"] .card-mnemonic-body .similar-hint-block .similar-title { color: #fbbf24; }
[data-theme="dark"] .card-mnemonic-body .similar-char-item {
    background: rgba(220, 38, 38, .12);
    border-color: rgba(248, 113, 113, .25);
}
[data-theme="dark"] .card-mnemonic-body .similar-char-item:hover {
    background: rgba(220, 38, 38, .2);
    border-color: rgba(248, 113, 113, .5);
}
[data-theme="dark"] .card-mnemonic-body .similar-char-btn-audio {
    background: rgba(167, 139, 250, .2);
    color: #c4b5fd;
}
[data-theme="dark"] .card-mnemonic-body .similar-char-btn-audio:hover {
    background: #a78bfa;
    color: #fff;
}
[data-theme="dark"] .card-mnemonic-body .similar-char-main .similar-char-zh { color: #fca5a5; }
[data-theme="dark"] .card-mnemonic-body .similar-char-main .similar-char-pinyin { color: #c4b5fd; }
[data-theme="dark"] .card-mnemonic-body .similar-char-btn-jump {
    background: rgba(220, 38, 38, .2);
    color: #fca5a5;
}
[data-theme="dark"] .card-mnemonic-body .similar-char-btn-jump:hover {
    background: #dc2626;
    color: #fff;
}

.card-mnemonic-body .similar-hint-block .similar-diff {
    margin-top: .5rem;
    padding-top: .5rem;
    border-top: 1px dashed rgba(220, 38, 38, .25);
    font-size: .72rem;
    color: var(--text-2);
    line-height: 1.5;
}
[data-theme="dark"] .card-mnemonic-body .similar-hint-block .similar-diff {
    border-top-color: rgba(248, 113, 113, .3);
    color: #fca5a5;
}

@media (max-width: 500px) {
    .card-mnemonic-body .similar-char-item {
        padding: .15rem .35rem .15rem .15rem;
        gap: .2rem;
    }
    .card-mnemonic-body .similar-char-btn-audio,
    .card-mnemonic-body .similar-char-btn-jump {
        width: 22px; height: 22px;
        font-size: .6rem;
    }
    .card-mnemonic-body .similar-char-main .similar-char-zh { font-size: .95rem; }
    .nav-back-btn {
        bottom: calc(12px + env(safe-area-inset-bottom));
        left: 12px;
        padding: .55rem .85rem;
        font-size: .78rem;
    }
    .pf-char-btn {
        font-size: 1.05rem;
        padding: .3rem .6rem;
    }
    .pf-char-info {
        padding: .7rem .85rem;
    }
    .pf-char-info-zh { font-size: 1.35rem; }
    .pf-char-info-pinyin { font-size: .9rem; }
    .pf-char-info-line { font-size: .78rem; }
}

body[data-vocab-mode="1"] .card {
    padding: clamp(.6rem, 1vw, .85rem) clamp(.7rem, 1.2vw, 1rem) !important;
    padding-left: clamp(.85rem, 1.4vw, 1.1rem) !important;
    margin-bottom: 0 !important;
    border-radius: 12px !important;
}
body[data-vocab-mode="1"] .card-header {
    margin-bottom: .35rem !important;
    padding-bottom: .35rem !important;
    gap: .35rem !important;
}
body[data-vocab-mode="1"] .card-tag.topic,
body[data-vocab-mode="1"] .card-tag.subject {
    display: none !important;
}
body[data-vocab-mode="1"] .card-tag.hsk {
    padding: .12rem .45rem !important;
    font-size: .62rem !important;
}
body[data-vocab-mode="1"] .card-stt {
    min-width: 24px !important;
    height: 24px !important;
    font-size: .62rem !important;
    padding: 0 .4rem !important;
}
body[data-vocab-mode="1"] .action-group { gap: .2rem !important; }
body[data-vocab-mode="1"] .action-group .audio-btn,
body[data-vocab-mode="1"] .action-group .write-btn,
body[data-vocab-mode="1"] .action-group .practice-full-btn,
body[data-vocab-mode="1"] .action-group .fav-btn {
    width: 26px !important;
    height: 26px !important;
    font-size: .7rem !important;
}
body[data-vocab-mode="1"] .card-body {
    margin-bottom: .4rem !important;
    padding-left: 0 !important;
}
body[data-vocab-mode="1"] .card-vi {
    font-size: .75rem !important;
    margin-bottom: .25rem !important;
    padding-left: 0 !important;
    color: var(--text-2) !important;
    font-weight: 600 !important;
}
body[data-vocab-mode="1"] .card-vi::before { display: none !important; }
body[data-vocab-mode="1"] .card-zh {
    font-size: clamp(1.5rem, 3.5vw, 2rem) !important;
    font-weight: 700 !important;
    line-height: 1.1 !important;
    margin-bottom: .2rem !important;
    display: inline-block !important;
    margin-right: .5rem !important;
}
body[data-vocab-mode="1"] .card-pinyin {
    font-size: .82rem !important;
    padding: .15rem .5rem !important;
    margin-bottom: 0 !important;
    display: inline-block !important;
    vertical-align: middle !important;
}
body[data-vocab-mode="1"] .card-radical {
    margin-top: .35rem !important;
    padding: .35rem .5rem !important;
    border-radius: 8px !important;
    flex-direction: row !important;
    align-items: center !important;
    gap: .4rem !important;
}
body[data-vocab-mode="1"] .card-radical-label { font-size: .6rem !important; margin: 0 !important; }
body[data-vocab-mode="1"] .card-radical-body { gap: .4rem !important; flex: 1 !important; min-width: 0 !important; }
body[data-vocab-mode="1"] .card-radical-box {
    width: 32px !important;
    height: 32px !important;
    font-size: 1rem !important;
    border-radius: 5px !important;
}
body[data-vocab-mode="1"] .card-radical-info { gap: .1rem !important; }
body[data-vocab-mode="1"] .card-radical-name { font-size: .72rem !important; line-height: 1.25 !important; }
body[data-vocab-mode="1"] .card-radical-meaning { font-size: .68rem !important; line-height: 1.3 !important; color: var(--text-3) !important; }
body[data-vocab-mode="1"] .card-mnemonic {
    margin-top: .35rem !important;
    padding: .4rem .6rem !important;
    border-radius: 8px !important;
    gap: .3rem !important;
}
body[data-vocab-mode="1"] .card-mnemonic-label { font-size: .6rem !important; letter-spacing: .3px !important; }
body[data-vocab-mode="1"] .card-mnemonic-body { font-size: .76rem !important; line-height: 1.5 !important; }
body[data-vocab-mode="1"] .card-mnemonic-body .char-zh { font-size: .9rem !important; padding: 0 .1rem !important; }
body[data-vocab-mode="1"] .similar-hint-block {
    margin-top: .4rem !important;
    padding: .4rem .55rem !important;
    border-radius: 8px !important;
    font-size: .72rem !important;
    line-height: 1.4 !important;
}
body[data-vocab-mode="1"] .similar-hint-block .similar-title { font-size: .62rem !important; margin-bottom: .3rem !important; }
body[data-vocab-mode="1"] .similar-hint-block .similar-chars { gap: .25rem .35rem !important; margin-bottom: 0 !important; }
body[data-vocab-mode="1"] .similar-char-item { padding: .12rem .3rem .12rem .15rem !important; gap: .2rem !important; border-radius: 6px !important; }
body[data-vocab-mode="1"] .similar-char-btn-audio,
body[data-vocab-mode="1"] .similar-char-btn-jump { width: 20px !important; height: 20px !important; font-size: .55rem !important; }
body[data-vocab-mode="1"] .similar-char-main { gap: .15rem !important; }
body[data-vocab-mode="1"] .similar-char-main .similar-char-zh { font-size: .88rem !important; }
body[data-vocab-mode="1"] .similar-char-main .similar-char-pinyin { font-size: .6rem !important; }
body[data-vocab-mode="1"] .similar-char-main .similar-char-info-txt { font-size: .6rem !important; }
body[data-vocab-mode="1"] .similar-sep { font-size: .7rem !important; }
body[data-vocab-mode="1"] .card-practice {
    padding-top: .35rem !important;
    margin-top: 0 !important;
    border-top: 1px dashed var(--border) !important;
    gap: .25rem !important;
}
body[data-vocab-mode="1"] .card-practice .practice-input { padding: .3rem .6rem !important; font-size: .78rem !important; }
body[data-vocab-mode="1"] .card-practice .toggle-check-btn { width: 26px !important; height: 26px !important; font-size: .7rem !important; }
@media (min-width: 769px) {
    body[data-vocab-mode="1"] .mobile-view {
        grid-template-columns: repeat(2, minmax(0, 1fr)) !important;
        gap: .6rem !important;
    }
}
@media (min-width: 1200px) {
    body[data-vocab-mode="1"] .mobile-view {
        grid-template-columns: repeat(2, minmax(0, 1fr)) !important;
        gap: .7rem !important;
    }
}
body[data-vocab-mode="1"] .card.focused { transform: scale(1.01) !important; }
body[data-vocab-mode="1"] .card.focused .card-zh { font-size: clamp(1.65rem, 3.8vw, 2.1rem) !important; }
@media (max-width: 500px) {
    body[data-vocab-mode="1"] .card {
        padding: .5rem .65rem !important;
        padding-left: .8rem !important;
        border-radius: 10px !important;
    }
    body[data-vocab-mode="1"] .card-header { margin-bottom: .25rem !important; padding-bottom: .25rem !important; }
    body[data-vocab-mode="1"] .card-zh { font-size: 1.35rem !important; }
    body[data-vocab-mode="1"] .card-pinyin { font-size: .74rem !important; }
    body[data-vocab-mode="1"] .card-radical-box { width: 28px !important; height: 28px !important; font-size: .9rem !important; }
    body[data-vocab-mode="1"] .card-mnemonic-body { font-size: .72rem !important; line-height: 1.45 !important; }
    body[data-vocab-mode="1"] .similar-char-item { padding: .1rem .25rem .1rem .12rem !important; }
    body[data-vocab-mode="1"] .similar-char-btn-audio,
    body[data-vocab-mode="1"] .similar-char-btn-jump { width: 18px !important; height: 18px !important; font-size: .5rem !important; }
}
body[data-vocab-mode="1"] { scroll-behavior: smooth; }
body[data-vocab-mode="1"] .mobile-view { gap: .5rem !important; }
@media (max-width: 500px) {
    body[data-vocab-mode="1"] .mobile-view { gap: .4rem !important; }
}

.pf-char-btn.reading {
    background: linear-gradient(135deg, #fbbf24, #f59e0b) !important;
    color: #fff !important;
    border-color: #f59e0b !important;
    box-shadow: 0 6px 20px rgba(245, 158, 11, .6), 0 0 0 4px rgba(245, 158, 11, .25) !important;
    transform: translateY(-3px) scale(1.15) !important;
    animation: charReadingPulse 1s ease-in-out infinite;
    position: relative;
    z-index: 10;
}
@keyframes charReadingPulse {
    0%, 100% {
        box-shadow: 0 6px 20px rgba(245, 158, 11, .6), 0 0 0 4px rgba(245, 158, 11, .25);
        transform: translateY(-3px) scale(1.15);
    }
    50% {
        box-shadow: 0 8px 28px rgba(245, 158, 11, .8), 0 0 0 8px rgba(245, 158, 11, .1);
        transform: translateY(-3px) scale(1.2);
    }
}
.pf-char-info-audio {
    width: 32px;
    height: 32px;
    border-radius: 50%;
    border: none;
    background: linear-gradient(135deg, #8b5cf6, #7c3aed);
    color: #fff;
    cursor: pointer;
    display: inline-flex;
    align-items: center;
    justify-content: center;
    font-size: .85rem;
    transition: all .2s;
    flex-shrink: 0;
    margin-left: auto;
    box-shadow: 0 3px 10px rgba(124, 58, 237, .4);
    padding: 0;
}
.pf-char-info-audio:hover {
    transform: scale(1.15);
    box-shadow: 0 5px 16px rgba(124, 58, 237, .6);
}
.pf-char-info-audio:active { transform: scale(.95); }
.pf-char-info-audio i { pointer-events: none; }
[data-theme="dark"] .pf-char-info-audio {
    background: linear-gradient(135deg, #a78bfa, #8b5cf6);
}
.pf-char-btn.zooming {
    animation: charZoomIn 0.6s cubic-bezier(.34, 1.56, .64, 1);
}
@keyframes charZoomIn {
    0%   { transform: scale(1); }
    40%  { transform: scale(1.35); }
    70%  { transform: scale(0.95); }
    100% { transform: scale(1); }
}
.pf-char-btn.active {
    transform: translateY(-2px) scale(1.08);
    box-shadow: 0 6px 18px rgba(124, 58, 237, .5);
}
.pf-karaoke-pinyin {
    padding: .35rem .75rem;
    background: linear-gradient(135deg, #fbbf24, #f59e0b);
    color: #fff;
    font-size: .82rem;
    font-weight: 800;
    font-family: inherit;
    font-style: italic;
    border-radius: 50px;
    box-shadow: 0 6px 20px rgba(245, 158, 11, .6), 0 0 0 3px rgba(255, 255, 255, .5);
    pointer-events: none;
    animation: karaokePinyinIn .25s cubic-bezier(.34, 1.56, .64, 1);
    white-space: nowrap;
    letter-spacing: .02em;
    text-shadow: 0 1px 2px rgba(0, 0, 0, .2);
}
.pf-karaoke-pinyin::after {
    content: '';
    position: absolute;
    top: 100%;
    left: 50%;
    transform: translateX(-50%);
    border: 6px solid transparent;
    border-top-color: #f59e0b;
    filter: drop-shadow(0 2px 2px rgba(245, 158, 11, .3));
}
@keyframes karaokePinyinIn {
    from {
        opacity: 0;
        transform: translate(-50%, -100%) scale(.8) translateY(8px);
    }
    to {
        opacity: 1;
        transform: translate(-50%, -100%) scale(1) translateY(0);
    }
}

.card-mnemonic-body .mnemonic-story-wrap{display:block;margin:.55rem 0 .1rem;border-radius:10px;overflow:hidden;background:linear-gradient(135deg,rgba(99,102,241,.07) 0%,rgba(139,92,246,.05) 100%);border:1px solid rgba(139,92,246,.18);transition:border-color .25s ease,box-shadow .25s ease}
.card-mnemonic-body .mnemonic-story-wrap.open{border-color:rgba(139,92,246,.4);box-shadow:0 4px 14px rgba(139,92,246,.12)}
.card-mnemonic-body .mnemonic-story-toggle{display:flex;align-items:center;justify-content:space-between;width:100%;padding:.55rem .75rem;background:transparent;border:none;font-family:inherit;font-size:.82rem;font-weight:800;color:#6d28d9;cursor:pointer;user-select:none;-webkit-tap-highlight-color:transparent;text-align:left;letter-spacing:.02em;transition:background .18s ease,color .18s ease;gap:.5rem}
.card-mnemonic-body .mnemonic-story-toggle:hover{background:rgba(139,92,246,.08);color:#5b21b6}
.card-mnemonic-body .mnemonic-story-toggle:active{background:rgba(139,92,246,.14)}
.card-mnemonic-body .mnemonic-story-toggle-label{display:inline-flex;align-items:center;gap:.4rem;flex:1;min-width:0}
.card-mnemonic-body .mnemonic-story-toggle-arrow{flex-shrink:0;display:inline-flex;align-items:center;justify-content:center;width:22px;height:22px;border-radius:50%;background:rgba(139,92,246,.15);color:#7c3aed;font-size:.7rem;transition:transform .3s cubic-bezier(.34,1.56,.64,1),background .2s ease}
.card-mnemonic-body .mnemonic-story-wrap.open .mnemonic-story-toggle-arrow{transform:rotate(180deg);background:#7c3aed;color:#fff}
.card-mnemonic-body .mnemonic-story-content{display:grid;grid-template-rows:0fr;transition:grid-template-rows .35s cubic-bezier(.4,0,.2,1);overflow:hidden}
.card-mnemonic-body .mnemonic-story-wrap.open .mnemonic-story-content{grid-template-rows:1fr}
.card-mnemonic-body .mnemonic-story-inner{overflow:hidden;padding:0 .75rem;font-size:.82rem;color:var(--text-2);line-height:1.6;transition:padding .3s ease}
.card-mnemonic-body .mnemonic-story-wrap.open .mnemonic-story-inner{padding:.1rem .75rem .7rem}
.card-mnemonic-body .mnemonic-story-inner .char-zh{font-family:var(--font-zh);font-size:1rem;font-weight:600;color:var(--text);padding:0 .15rem}
.card-mnemonic-body .mnemonic-story-inner .arrow{color:#f59e0b;font-weight:800;padding:0 .2rem}
.card-mnemonic-body .mnemonic-story-inner .hint{color:#b45309;font-weight:700}
[data-theme="dark"] .card-mnemonic-body .mnemonic-story-wrap{background:linear-gradient(135deg,rgba(139,92,246,.12) 0%,rgba(99,102,241,.08) 100%);border-color:rgba(167,139,250,.25)}
[data-theme="dark"] .card-mnemonic-body .mnemonic-story-wrap.open{border-color:rgba(167,139,250,.5);box-shadow:0 4px 14px rgba(139,92,246,.25)}
[data-theme="dark"] .card-mnemonic-body .mnemonic-story-toggle{color:#c4b5fd}
[data-theme="dark"] .card-mnemonic-body .mnemonic-story-toggle:hover{background:rgba(139,92,246,.15);color:#ddd6fe}
[data-theme="dark"] .card-mnemonic-body .mnemonic-story-toggle-arrow{background:rgba(167,139,250,.25);color:#c4b5fd}
[data-theme="dark"] .card-mnemonic-body .mnemonic-story-wrap.open .mnemonic-story-toggle-arrow{background:#a78bfa;color:#1e1b4b}
[data-theme="dark"] .card-mnemonic-body .mnemonic-story-inner{color:#cbd5e1}
[data-theme="dark"] .card-mnemonic-body .mnemonic-story-inner .char-zh{color:#e0f2fe}
body[data-vocab-mode="1"] .card-mnemonic-body .mnemonic-story-wrap{margin:.35rem 0 .1rem!important;border-radius:8px!important}
body[data-vocab-mode="1"] .card-mnemonic-body .mnemonic-story-toggle{padding:.4rem .55rem!important;font-size:.72rem!important}
body[data-vocab-mode="1"] .card-mnemonic-body .mnemonic-story-toggle-arrow{width:18px!important;height:18px!important;font-size:.6rem!important}
body[data-vocab-mode="1"] .card-mnemonic-body .mnemonic-story-inner{font-size:.74rem!important;line-height:1.5!important}
body[data-vocab-mode="1"] .card-mnemonic-body .mnemonic-story-wrap.open .mnemonic-story-inner{padding:.05rem .55rem .55rem!important}
"""
    return css.replace("__VOCAB_ID__", vocab_id)


def build_vocab_tab_html(vocab_id="tu-vung", label="Từ vựng HSK"):
    return (
        '\n        <button class="ds-btn ds-btn-primary" '
        'data-dataset="' + vocab_id + '" id="dsTuVungBtn">\n'
        '            <i class="fas fa-book"></i>\n'
        '            <span id="dsTuVungLabel">' + label + '</span>\n'
        '            <span class="ds-vocab-badge" id="dsVocabBadge">PREMIUM</span>\n'
        '        </button>'
    )


def build_vocab_modal_html():
    return r'''
<div class="vocab-upgrade-modal" id="vocabUpgradeModal">
    <div class="vocab-upgrade-box">
        <div class="vocab-upgrade-header">
            <div class="vocab-upgrade-icon" id="vocabUpgradeIcon">👑</div>
            <div class="vocab-upgrade-title" id="vocabUpgradeTitle">Mở khóa Từ vựng HSK</div>
            <div class="vocab-upgrade-subtitle" id="vocabUpgradeSubtitle">
                Gia hạn gói để sử dụng Từ vựng HSK và toàn bộ tính năng
            </div>
            <div class="vocab-upgrade-price" id="vocabUpgradePrice" style="display:none;">
                💎 1.000.000đ
            </div>
        </div>
        <div class="vocab-upgrade-body">
            <div class="vocab-upgrade-features">
                <div class="vocab-upgrade-feature">
                    <i class="fas fa-book"></i>
                    <span>Toàn bộ từ vựng HSK 1–9 (~11700 từ)</span>
                </div>
                <div class="vocab-upgrade-feature">
                    <i class="fas fa-lightbulb"></i>
                    <span>Mẹo nhớ chữ Hán chi tiết</span>
                </div>
                <div class="vocab-upgrade-feature">
                    <i class="fas fa-pen-fancy"></i>
                    <span>Phân tích bộ thủ từng chữ</span>
                </div>
                <div class="vocab-upgrade-feature">
                    <i class="fas fa-quote-left"></i>
                    <span>Câu ví dụ + phiên âm + nghĩa</span>
                </div>
                <div class="vocab-upgrade-feature highlight" id="vocabUpgradeFeatureHighlight">
                    <i class="fas fa-crown"></i>
                    <span><b>Nhiều gói linh hoạt — Chọn gói phù hợp</b></span>
                </div>
            </div>
            <div class="vocab-upgrade-actions" id="vocabUpgradeActions"></div>
        </div>
    </div>
</div>
'''


def build_vocab_js_override(vocab_id="tu-vung"):
    load_similar_chars()
    similar_json = json.dumps(_SIMILAR_CHARS, ensure_ascii=False)

    js = r"""
window.__SIMILAR_CHARS__ = __SIMILAR_JSON__;

(function() {
    'use strict';

    var VOCAB_ID = '__VOCAB_ID__';
    var _done = new WeakSet();

    var _navStack = [];
    var _navBackBtn = null;

    function canAccessVocab() { return true; }

    function _esc(s) {
        return (typeof escapeHtml === 'function')
            ? escapeHtml(s) : String(s == null ? '' : s);
    }

    function _isVocabMode() {
        return (typeof CURRENT_DATASET !== 'undefined') && CURRENT_DATASET === VOCAB_ID;
    }

    function _findRecord(stt) {
        if (stt == null) return null;
        var s = String(stt);
        if (window.FIXPY_DATASETS && window.FIXPY_DATASETS[VOCAB_ID]) {
            var list = window.FIXPY_DATASETS[VOCAB_ID].data || [];
            for (var i = 0; i < list.length; i++) {
                if (String(list[i].stt) === s) return list[i];
            }
        }
        if (typeof RAW_DATA !== 'undefined' && RAW_DATA) {
            for (var j = 0; j < RAW_DATA.length; j++) {
                if (String(RAW_DATA[j].stt) === s) return RAW_DATA[j];
            }
        }
        return null;
    }

    function _findCardByChar(char) {
        if (!char) return null;
        var cards = document.querySelectorAll('.card[data-stt]');
        for (var i = 0; i < cards.length; i++) {
            var zhEl = cards[i].querySelector('.card-zh');
            if (zhEl) {
                var text = (zhEl.textContent || '').trim();
                if (text === char) return cards[i];
            }
        }
        for (var j = 0; j < cards.length; j++) {
            var zhEl2 = cards[j].querySelector('.card-zh');
            if (zhEl2) {
                var text2 = (zhEl2.textContent || '').trim();
                if (text2.indexOf(char) !== -1) return cards[j];
            }
        }
        return null;
    }

    function _hasCardForChar(char) {
        if (!char) return false;
        return _findCardByChar(char) !== null;
    }

    function _scrollToCard(card) {
        if (!card) return;
        var prevScroll = window.pageYOffset
                      || document.documentElement.scrollTop
                      || document.body.scrollTop
                      || window.scrollY
                      || 0;
        _navStack.push(prevScroll);
        if (document.body.style.overflow === 'hidden') {
            document.body.style.overflow = '';
        }
        var rect = card.getBoundingClientRect();
        var yOffset = rect.top + window.pageYOffset - 100;
        try {
            window.scrollTo({ top: yOffset, behavior: 'smooth' });
        } catch (e) {
            window.scrollTo(0, yOffset);
        }
        card.classList.add('nav-highlight');
        setTimeout(function() {
            card.classList.remove('nav-highlight');
        }, 2000);
        _showNavBackBtn();
    }

    function _showNavBackBtn() {
        if (_navBackBtn) _navBackBtn.remove();
        _navBackBtn = document.createElement('button');
        _navBackBtn.className = 'nav-back-btn';
        _navBackBtn.type = 'button';
        _navBackBtn.innerHTML = '<i class="fas fa-arrow-left"></i> Quay lại';
        _navBackBtn.addEventListener('click', function(e) {
            e.stopImmediatePropagation();
            e.stopPropagation();
            e.preventDefault();
            _navBack();
        }, true);
        document.body.appendChild(_navBackBtn);
    }

    function _navBack() {
        if (_navStack.length === 0) {
            if (_navBackBtn) _navBackBtn.remove();
            _navBackBtn = null;
            return;
        }
        var prevScroll = _navStack.pop();
        if (document.body.style.overflow === 'hidden') {
            document.body.style.overflow = '';
        }
        try {
            window.scrollTo({ top: prevScroll, behavior: 'smooth' });
        } catch (e) {
            window.scrollTo(0, prevScroll);
        }
        setTimeout(function() {
            var currentScroll = window.pageYOffset
                             || document.documentElement.scrollTop
                             || document.body.scrollTop
                             || 0;
            if (Math.abs(currentScroll - prevScroll) > 50) {
                document.documentElement.scrollTop = prevScroll;
                document.body.scrollTop = prevScroll;
            }
        }, 500);
        if (_navStack.length === 0) {
            setTimeout(function() {
                if (_navBackBtn) _navBackBtn.remove();
                _navBackBtn = null;
            }, 300);
        }
    }

    function _hideNavBackBtn() {
        if (_navBackBtn) {
            _navBackBtn.remove();
            _navBackBtn = null;
        }
        _navStack = [];
    }

    var _lastDatasetForNav = null;
    function watchDatasetNav() {
        var curDataset = (typeof CURRENT_DATASET !== 'undefined') ? CURRENT_DATASET : '';
        if (_lastDatasetForNav === null) {
            _lastDatasetForNav = curDataset;
            return;
        }
        if (_lastDatasetForNav === VOCAB_ID && curDataset !== VOCAB_ID) {
            _hideNavBackBtn();
        }
        _lastDatasetForNav = curDataset;
    }

    window.vocabSpeakChar = function(text, btn, evt) {
        if (evt) { evt.stopPropagation(); evt.preventDefault(); }
        if (!text) return;
        if (!('speechSynthesis' in window)) return;
        speechSynthesis.cancel();
        if (btn) {
            document.querySelectorAll('.similar-char-btn-audio.speaking').forEach(function(b) {
                b.classList.remove('speaking');
            });
            document.querySelectorAll('.card-example-zh.speaking, .card-char-speakable.speaking').forEach(function(b) {
                b.classList.remove('speaking');
            });
            btn.classList.add('speaking');
        }
        var u = new SpeechSynthesisUtterance(text);
        u.lang = 'zh-CN';
        u.rate = text.length > 3 ? 0.75 : 0.85;
        if (typeof applyVoiceSettings === 'function') {
            try { applyVoiceSettings(u); } catch(e) {}
        }
        u.onend = u.onerror = function() {
            if (btn) btn.classList.remove('speaking');
        };
        setTimeout(function() { speechSynthesis.speak(u); }, 30);
    };
    window.vocabToggleStory = function(btn, evt) {
    if (evt) { evt.stopPropagation(); evt.preventDefault(); }
    if (!btn) return;
    var wrap = btn.closest('.mnemonic-story-wrap');
    if (!wrap) return;
    var isOpen = wrap.classList.toggle('open');
    btn.setAttribute('aria-expanded', isOpen ? 'true' : 'false');
};

    window.vocabJumpToChar = function(char, btn, evt) {
        if (evt) { evt.stopPropagation(); evt.preventDefault(); }
        if (!char) return;
        var targetCard = _findCardByChar(char);
        if (targetCard) {
            _scrollToCard(targetCard);
        } else {
            if (typeof showTagToast === 'function') {
                showTagToast('Không tìm thấy thẻ "' + char + '"');
            }
        }
    };

    function updateTabLockState() {
        var btn = document.querySelector('.ds-btn[data-dataset="' + VOCAB_ID + '"]');
        if (!btn) return;
        var oldLock = btn.querySelector('.vocab-lock-icon');
        if (oldLock) oldLock.remove();
        btn.classList.remove('vocab-locked');
        btn.title = 'Từ vựng HSK - đã mở khóa';
    }

    function openUpgradeModal() {
        var modal = document.getElementById('vocabUpgradeModal');
        if (!modal) return;
        var titleEl = document.getElementById('vocabUpgradeTitle');
        var subEl = document.getElementById('vocabUpgradeSubtitle');
        var actionsEl = document.getElementById('vocabUpgradeActions');
        var iconEl = document.getElementById('vocabUpgradeIcon');
        var highlightEl = document.getElementById('vocabUpgradeFeatureHighlight');

        if (iconEl) iconEl.textContent = '🔐';
        if (titleEl) titleEl.textContent = 'Đăng nhập để sử dụng Từ vựng HSK';
        if (subEl) subEl.textContent = 'Đăng nhập miễn phí để trải nghiệm kho từ vựng HSK 1-9';
        if (highlightEl) {
            highlightEl.innerHTML =
                '<i class="fas fa-gift"></i>' +
                '<span><b>Đăng nhập miễn phí — Dùng thử ngay</b></span>';
        }
        if (actionsEl) {
            actionsEl.innerHTML =
                '<button class="vocab-upgrade-btn primary" onclick="vocabUpgradeLogin()">' +
                    '<i class="fas fa-sign-in-alt"></i> Đăng nhập' +
                '</button>' +
                '<button class="vocab-upgrade-btn secondary" onclick="vocabUpgradeClose()">' +
                    'Để sau' +
                '</button>';
        }
        modal.classList.add('show');
        document.body.style.overflow = 'hidden';
    }

    function closeUpgradeModal() {
        var modal = document.getElementById('vocabUpgradeModal');
        if (modal) modal.classList.remove('show');
        document.body.style.overflow = '';
    }

    window.vocabUpgradeClose = closeUpgradeModal;
    window.vocabUpgradeLogin = function() {
        closeUpgradeModal();
        if (typeof showLoginModal === 'function') showLoginModal();
    };
    window.vocabUpgradeRenew = function() {
        closeUpgradeModal();
        document.body.style.overflow = '';
        if (typeof openRenewalModal === 'function') {
            setTimeout(function() { openRenewalModal(); }, 150);
        }
    };

    document.addEventListener('click', function(e) {
        var modal = document.getElementById('vocabUpgradeModal');
        if (!modal || !modal.classList.contains('show')) return;
        if (e.target === modal) closeUpgradeModal();
    });
    document.addEventListener('keydown', function(e) {
        if (e.key === 'Escape') {
            var modal = document.getElementById('vocabUpgradeModal');
            if (modal && modal.classList.contains('show')) closeUpgradeModal();
        }
    });

    function highlightPinyin(text) {
        if (!text) return '';
        return text.replace(
            /[a-zA-ZàáảãạăằắẳẵặâầấẩẫậèéẻẽẹêềếểễệìíỉĩịòóỏõọôồốổỗộơờớởỡợùúủũụưừứửữựỳýỷỹỵđÀÁẢÃẠĂẰẮẲẴẶÂẦẤẨẪẬÈÉẺẼẸÊỀẾỂỄỆÌÍỈĨỊÒÓỎÕỌÔỒỐỔỖỘƠỜỚỞỠỢÙÚỦŨỤƯỪỨỬỮỰỲÝỶỸỴĐāáǎàēéěèīíǐìōóǒòūúǔùǖǘǚǜ]+/g,
            function(w) { return '<span class="pinyin-hl">' + w + '</span>'; }
        );
    }

    function buildRadicalBlock(radical) {
        if (!radical) return '';
        var zh = _esc(radical.zh || '');
        var py = _esc(radical.pinyin || '');
        var st = _esc(radical.strokes || '');
        var mean = _esc(radical.meaning || '');

        var nameLine = '';
        if (zh) {
            nameLine = '<span class="char">' + zh + '</span>';
            if (py) nameLine += ' <span class="pinyin">(' + py + ')</span>';
            if (st) nameLine += ' - ' + st + ' net';
        }

        var html = '<div class="card-radical">';
        html += '<div class="card-radical-label">BỘ THỦ</div>';
        html += '<div class="card-radical-body">';
        if (zh) html += '<div class="card-radical-box">' + zh + '</div>';
        html += '<div class="card-radical-info">';
        if (nameLine) html += '<div class="card-radical-name">' + nameLine + '</div>';
        if (mean) html += '<div class="card-radical-meaning">' + mean + '</div>';
        html += '</div></div></div>';
        return html;
    }

    function buildMnemonicBlock(text, currentChar, viDuList) {
    currentChar = currentChar || '';
    if (!text || !text.trim()) return '';

    var mnemonicPart = text;
    var storyPart = '';

    var storyMatch = text.match(
        /🎬\s*([\s\S]*?)(?=(?:\n\s*📎\s*Ví\s*dụ)|\n\s*🔗|\n\s*💡|\n\s*📌|\s*$)/
    );

    if (storyMatch) {
        storyPart = storyMatch[0].trim();
        mnemonicPart = text.replace(storyPart, '').replace(/\n{2,}/g, '\n').trim();
    }

    function _processText(rawIn) {
        if (!rawIn) return '';

        var raw = rawIn;
        raw = raw.replace(/\s=\s/g, ' __ARROW__ ');
        raw = raw.replace(/→/g, ' __RARR__ ');
        raw = raw.replace(/\(([^)]+)\)/g, ' __PAREN_L__$1__PAREN_R__');

        var safe = _esc(raw);

        function _normZh(s) {
            if (!s) return '';
            return s.replace(/[\s，。！？、；：""''（）]/g, '').trim();
        }

        var viDuNormMap = {};
        if (viDuList && viDuList.length > 0) {
            viDuList.forEach(function(ex) {
                if (ex && ex.zh) {
                    viDuNormMap[_normZh(ex.zh)] = ex;
                }
            });
        }

        safe = safe.replace(
            /([\u4e00-\u9fff][\u4e00-\u9fff\s]*[\u4e00-\u9fff，。！？、；：""''（）])/g,
            function(match) {
                var trimmed = match.trim();
                var hanziOnly = trimmed.replace(/[^\u4e00-\u9fff]/g, '');
                if (hanziOnly.length < 2) return match;

                var mJs = trimmed
                    .replace(/\\/g, '\\\\')
                    .replace(/'/g, "\\'")
                    .replace(/"/g, '\\"');
                var mNorm = _normZh(trimmed);

                var isExample = !!viDuNormMap[mNorm];
                var cls = isExample ? 'card-example-zh' : 'card-char-speakable';
                var title = isExample ? 'Bấm để nghe câu ví dụ' : 'Bấm để nghe';

                return '<span class="' + cls + '" '
                     + 'onclick="vocabSpeakChar(\'' + mJs + '\', this, event)" '
                     + 'title="' + title + '">'
                     + match + '</span>';
            }
        );

        safe = safe.replace(/__ARROW__/g, '<span class="arrow">=</span>');
        safe = safe.replace(/__RARR__/g, '<span class="arrow">→</span>');
        safe = safe.replace(/__PAREN_L__/g, '(<span class="hint">');
        safe = safe.replace(/__PAREN_R__/g, '</span>)');

        return safe;
    }

    var mnemonicHtml = _processText(mnemonicPart);

    var storyHtml = '';
    if (storyPart) {
        var storyBody = storyPart.replace(/^🎬\s*/, '').trim();
        var storyContentHtml = _processText(storyBody);

        storyHtml =
            '<div class="mnemonic-story-wrap" data-story="1">' +
                '<button class="mnemonic-story-toggle" type="button" ' +
                        'onclick="vocabToggleStory(this, event)" ' +
                        'aria-expanded="false">' +
                    '<span class="mnemonic-story-toggle-label">' +
                        '<span>🎬</span>' +
                        '<span>Câu chuyện ghi nhớ</span>' +
                    '</span>' +
                    '<span class="mnemonic-story-toggle-arrow">' +
                        '<i class="fas fa-chevron-down"></i>' +
                    '</span>' +
                '</button>' +
                '<div class="mnemonic-story-content">' +
                    '<div class="mnemonic-story-inner">' +
                        storyContentHtml +
                    '</div>' +
                '</div>' +
            '</div>';
    }

    return '<div class="card-mnemonic">'
        + '<div class="card-mnemonic-label">MẸO NHỚ</div>'
        + '<div class="card-mnemonic-body">'
        + mnemonicHtml
        + storyHtml
        + '</div>'
        + '</div>';
}
    function buildSimilarCharsBlock(currentChar) {
        currentChar = currentChar || '';
        if (!currentChar) return '';

        var similarMap = window.__SIMILAR_CHARS__ || {};
        var entry = similarMap[currentChar];

        var chars = [];
        if (entry && typeof entry === 'object') {
            if (Array.isArray(entry)) {
                chars = entry;
            } else if (Array.isArray(entry.similar)) {
                chars = entry.similar;
            }
        }

        if (!chars || chars.length === 0) return '';

        var vocabList = [];
        if (window.FIXPY_DATASETS && window.FIXPY_DATASETS[VOCAB_ID]) {
            vocabList = window.FIXPY_DATASETS[VOCAB_ID].data || [];
        }
        if (vocabList.length === 0 && typeof RAW_DATA !== 'undefined' && RAW_DATA) {
            vocabList = RAW_DATA;
        }

        var items = [];
        for (var i = 0; i < chars.length; i++) {
            var c = String(chars[i] || '').trim();
            if (!c) continue;

            var info = null;
            for (var j = 0; j < vocabList.length; j++) {
                if (vocabList[j].zh === c) { info = vocabList[j]; break; }
            }
            if (!info) {
                for (var k = 0; k < vocabList.length; k++) {
                    if (vocabList[k].zh && vocabList[k].zh.indexOf(c) !== -1) {
                        info = vocabList[k];
                        break;
                    }
                }
            }

            var pinyin = info ? (info.pinyin || '') : '';
            var vi = info ? (info.vi || '') : '';
            if (vi.length > 20) vi = vi.substring(0, 20) + '...';

            var extra = '';
            if (pinyin && vi) {
                extra = '<span class="similar-char-info-txt">'
                      + '<span class="similar-char-pinyin">/' + _esc(pinyin) + '/</span> '
                      + _esc(vi)
                      + '</span>';
            } else if (pinyin) {
                extra = '<span class="similar-char-info-txt">'
                      + '<span class="similar-char-pinyin">/' + _esc(pinyin) + '/</span>'
                      + '</span>';
            } else if (vi) {
                extra = '<span class="similar-char-info-txt">' + _esc(vi) + '</span>';
            } else {
                extra = '<span class="similar-char-info-txt" style="opacity:.5">(?)</span>';
            }

            var charJs = c.replace(/\\/g, '\\\\').replace(/'/g, "\\'").replace(/"/g, '\\"');
            var hasCard = _hasCardForChar(c);

            var jumpBtn = '';
            if (hasCard) {
                jumpBtn = '<button class="similar-char-btn-jump" '
                        + 'onclick="vocabJumpToChar(\'' + charJs + '\', this, event)" '
                        + 'title="Xem chi tiết">'
                        + '<i class="fas fa-arrow-right"></i></button>';
            }

            items.push(
                '<span class="similar-char-item" data-char="' + _esc(c) + '">'
                + '<button class="similar-char-btn-audio" '
                + 'onclick="vocabSpeakChar(\'' + charJs + '\', this, event)" '
                + 'title="Đọc âm">'
                + '<i class="fas fa-volume-up"></i></button>'
                + '<span class="similar-char-main">'
                + '<span class="similar-char-zh">' + _esc(c) + '</span>'
                + extra
                + '</span>'
                + jumpBtn
                + '</span>'
            );
        }

        if (items.length === 0) return '';

        var diffText = '';
        if (entry && !Array.isArray(entry) && entry.diff) {
            diffText = '<div class="similar-diff">'
                     + '📌 ' + _esc(entry.diff)
                     + '</div>';
        }

        return '<div class="similar-hint-block">'
             + '<span class="similar-title">🔍 DỄ NHẦM</span>'
             + '<div class="similar-chars">'
             + items.join(' <span class="similar-sep">·</span> ')
             + '</div>'
             + diffText
             + '</div>';
    }

    function enhanceCards() {
        if (!_isVocabMode()) return;
        if (!canAccessVocab()) return;

        var cards = document.querySelectorAll('.card[data-stt]');
        cards.forEach(function(card) {
            if (_done.has(card)) return;
            _done.add(card);

            var r = _findRecord(card.dataset.stt);
            if (!r) return;

            var body = card.querySelector('.card-body');
            if (!body) return;

            var oldR = body.querySelector('.card-radical');
            if (oldR) oldR.remove();
            var oldM = body.querySelector('.card-mnemonic');
            if (oldM) oldM.remove();

            var sttEl = card.querySelector('.card-stt');
            if (sttEl && r.stt_original) {
                sttEl.textContent = r.stt_original;
            }

            var anchor = body.querySelector('.card-vocab-example');
            if (r.radical) {
                var htmlR = buildRadicalBlock(r.radical);
                if (anchor) anchor.insertAdjacentHTML('beforebegin', htmlR);
                else body.insertAdjacentHTML('beforeend', htmlR);
            }

            if (r.mnemonic) {
                body.insertAdjacentHTML('beforeend', buildMnemonicBlock(r.mnemonic, r.zh, r.vi_du_list));
            }

            var similarHtml = buildSimilarCharsBlock(r.zh);
            if (similarHtml) {
                var mnemonicBody = body.querySelector('.card-mnemonic-body');
                if (mnemonicBody) {
                    mnemonicBody.insertAdjacentHTML('beforeend', similarHtml);
                } else {
                    body.insertAdjacentHTML('beforeend', similarHtml);
                }
            }

            var pyEl = body.querySelector('.card-vocab-example-pinyin')
                    || body.querySelector('.card-pinyin');
            if (pyEl && !pyEl.dataset.hlDone) {
                pyEl.dataset.hlDone = '1';
                pyEl.innerHTML = highlightPinyin(pyEl.textContent);
            }
        });
    }

    function setupObserver() {
        var wrapper = document.getElementById('mobileWrapper');
        if (!wrapper) { setTimeout(setupObserver, 400); return; }
        var timer = null;
        var observer = new MutationObserver(function() {
            if (timer) clearTimeout(timer);
            timer = setTimeout(enhanceCards, 80);
        });
        observer.observe(wrapper, { childList: true, subtree: true });
        enhanceCards();
    }

    function bindTabIfNeeded() {
        var btn = document.querySelector('.ds-btn[data-dataset="' + VOCAB_ID + '"]');
        if (!btn || btn.__vocabPremiumBound) return;

        btn.__vocabPremiumBound = true;

        btn.addEventListener('click', function(e) {
            e.stopImmediatePropagation();
            e.stopPropagation();
            e.preventDefault();

            window.__onboardingOverride = null;
            window.__onboardingAutoPicked = false;

            if (typeof state !== 'undefined' && state) {
                state.search = '';
                state.hsk = '';
                state.subject = '';
            }

            try {
                var si = document.getElementById('searchInput');
                var hf = document.getElementById('hskFilter');
                var sf = document.getElementById('subjectFilter');
                if (si) si.value = '';
                if (hf) hf.value = '';
                if (sf) sf.value = '';
                var cb = document.getElementById('clearSearchBtn');
                if (cb) cb.classList.remove('show');
            } catch(err) {}

            var subWrap = document.getElementById('dsSubWrap');
            if (subWrap) {
                subWrap.style.display = 'none';
                subWrap.classList.remove('show');
            }

            var oldOnbBanner = document.getElementById('onboardingActiveBanner');
            if (oldOnbBanner) oldOnbBanner.remove();
            var oldAnyBanner = document.querySelectorAll('[id*="onboarding"][id*="banner"]');
            oldAnyBanner.forEach(function(b) {
                if (b.id !== 'vocabWarningBanner') b.remove();
            });

            try {
                if (typeof buildFilters === 'function') buildFilters();
            } catch(err) {}

            document.querySelectorAll('.ds-btn, .ds-sub-btn').forEach(function(b) {
                b.classList.remove('active');
            });
            this.classList.add('active');

            if (typeof window.__switchRawData === 'function') {
                window.__switchRawData(VOCAB_ID);
            }

            try {
                if (typeof applyFilter === 'function') applyFilter();
                if (typeof updateResultCount === 'function') updateResultCount();
            } catch(err) {}

            setTimeout(function() {
                var mainEl = document.getElementById('mainContent');
                if (mainEl) {
                    var yOffset = mainEl.getBoundingClientRect().top + window.scrollY - 100;
                    window.scrollTo({ top: yOffset, behavior: 'smooth' });
                }
            }, 100);

            setTimeout(function() {
                if (typeof enhanceCards === 'function') enhanceCards();
            }, 300);

        }, true);
    }

    function hookPracticeFull() {
        if (typeof window.loadPracticeFull !== 'function') return false;
        if (window.loadPracticeFull.__vocabHooked) return true;

        var orig = window.loadPracticeFull;
        window.loadPracticeFull = function(stt) {
            var result = orig.apply(this, arguments);

            if (_isVocabMode() && canAccessVocab()) {
                var r = _findRecord(stt);
                if (r && r.vi_du_zh) {
                    try {
                        if (typeof window.pfCurrentAnswer !== 'undefined') {
                            window.pfCurrentAnswer = r.vi_du_zh;
                        }
                        if (typeof window.pfCurrentVi !== 'undefined') {
                            window.pfCurrentVi = r.vi_du_vi;
                        }
                        if (typeof window.pfCurrentPinyin !== 'undefined') {
                            window.pfCurrentPinyin = r.vi_du_pinyin;
                        }
                    } catch(e) {}

                    var pfViEl = document.getElementById('pfVi');
                    if (pfViEl && r.vi_du_vi) {
                        pfViEl.textContent = r.vi_du_vi;
                    }

                    setTimeout(function() {
                        var pfInput = document.getElementById('pfInput');
                        if (pfInput) pfInput.value = '';

                        var pfPreview = document.getElementById('pfPreview');
                        if (pfPreview) pfPreview.innerHTML = '';

                        var pfStatus = document.getElementById('pfStatus');
                        if (pfStatus) {
                            pfStatus.textContent = '';
                            pfStatus.className = 'practice-full-status';
                        }

                        var oldExample = document.querySelector('.pf-vocab-example');
                        if (oldExample) oldExample.remove();

                        var oldInfo = document.getElementById('pfCharInfo');
                        if (oldInfo) {
                            oldInfo.style.display = 'none';
                            oldInfo.innerHTML = '';
                        }

                        var hintBtn = document.getElementById('pfHintBtn');
                        if (hintBtn) {
                            hintBtn.classList.remove('active');
                            hintBtn.innerHTML = '<i class="fas fa-lightbulb"></i> Gợi ý';
                        }

                        if (hintBtn && !hintBtn.__vocabHooked) {
                            hintBtn.__vocabHooked = true;
                            hintBtn.addEventListener('click', function(ev) {
                                setTimeout(function() {
                                    if (!_isVocabMode()) return;
                                    if (!canAccessVocab()) return;

                                    var oldEx = document.querySelector('.pf-vocab-example');
                                    if (oldEx) oldEx.remove();

                                    var oldInfo2 = document.getElementById('pfCharInfo');
                                    if (oldInfo2) {
                                        oldInfo2.style.display = 'none';
                                        oldInfo2.innerHTML = '';
                                    }

                                    if (!hintBtn.classList.contains('active')) return;

                                    var currentStt = null;
                                    try { currentStt = window.pfCurrentStt; } catch(e) {}
                                    if (!currentStt) return;

                                    var rec2 = _findRecord(currentStt);
                                    if (!rec2 || !rec2.vi_du_zh) return;

                                    var answerEl2 = document.getElementById('pfAnswer');
                                    if (!answerEl2) return;

                                    var html2 = buildPFExampleBlock(rec2);
                                    if (html2) {
                                        answerEl2.insertAdjacentHTML('afterend', html2);
                                    }
                                }, 100);
                            });
                        }
                    }, 100);
                }
            }

            return result;
        };
        window.loadPracticeFull.__vocabHooked = true;
        return true;
    }

    function hookPfBuildQuickNav() {
        if (typeof window.pfBuildQuickNav !== 'function') return;
        if (window.pfBuildQuickNav.__vocabLimited) return;

        var orig = window.pfBuildQuickNav;
        window.pfBuildQuickNav = function() {
            if (!_isVocabMode()) {
                return orig.apply(this, arguments);
            }

            var sel = document.getElementById('pfQuickNav');
            if (!sel) return;

            var MAX_OPTIONS = 500;
            var list = (typeof filtered !== 'undefined') ? filtered : [];

            var total = list.length;
            var limit = Math.min(total, MAX_OPTIONS);

            var curStt = (typeof pfCurrentStt !== 'undefined' && pfCurrentStt)
                         ? String(pfCurrentStt)
                         : '';

            var curIdx = -1;
            for (var k = 0; k < list.length; k++) {
                if (String(list[k].stt) === curStt) {
                    curIdx = k;
                    break;
                }
            }

            var html = '<option value="">-- Chọn câu (' + total + ') --</option>';

            if (curIdx >= MAX_OPTIONS && curStt) {
                var rCur = list[curIdx];
                var viCur = (rCur.vi || '').substring(0, 45);
                var sttDisplayCur = rCur.stt_original || rCur.stt;
                var sttRawCur = (sttDisplayCur !== undefined
                              && sttDisplayCur !== null
                              && String(sttDisplayCur).trim() !== '')
                                ? '#' + String(sttDisplayCur).trim() + ' · '
                                : '';
                var labelCur = '⭐ ' + sttRawCur + 'Câu ' + (curIdx + 1) + ': ' + viCur;

                html += '<option value="' + _esc(rCur.stt) + '" selected>'
                      + _esc(labelCur)
                      + '</option>';
                html += '<option value="" disabled>───────────────</option>';
            }

            for (var i = 0; i < limit; i++) {
                var r = list[i];
                var vi = (r.vi || '').substring(0, 45);
                var sttDisplay = r.stt_original || r.stt;
                var sttRaw = (sttDisplay !== undefined
                           && sttDisplay !== null
                           && String(sttDisplay).trim() !== '')
                             ? '#' + String(sttDisplay).trim() + ' · '
                             : '';
                var label = sttRaw + 'Câu ' + (i + 1) + ': ' + vi;

                var isSelected = (String(r.stt) === curStt) ? ' selected' : '';

                html += '<option value="' + _esc(r.stt) + '"' + isSelected + '>'
                      + _esc(label)
                      + '</option>';
            }

            if (total > MAX_OPTIONS) {
                html += '<option value="" disabled>'
                      + '── Còn ' + (total - MAX_OPTIONS) + ' câu nữa, dùng nút ▶ ──'
                      + '</option>';
            }

            sel.innerHTML = html;

            if (curIdx >= MAX_OPTIONS && curStt) {
                sel.selectedIndex = 1;
            } else if (curStt) {
                try {
                    sel.value = curStt;
                } catch(e) {}
            }
        };

        window.pfBuildQuickNav.__vocabLimited = true;
    }

    function buildPFExampleBlock(r) {
        if (r && r.vi_du_zh) {
            window.__currentVocabExampleZh = r.vi_du_zh;
        }
        if (r && r.vi_du_pinyin) {
            window.__currentVocabExamplePinyin = r.vi_du_pinyin;
        }

        var zh = _esc(r.vi_du_zh || '');
        if (!zh) return '';

        var zhJs = String(r.vi_du_zh || '')
            .replace(/\\/g, '\\\\')
            .replace(/'/g, "\\'")
            .replace(/"/g, '\\"');
        var pinyin = _esc(r.vi_du_pinyin || '');
        var vi = _esc(r.vi_du_vi || '');

        var chars = (r && Array.isArray(r.vi_du_words)) ? r.vi_du_words : [];
        if (chars.length === 0) {
            for (var i = 0; i < zh.length; i++) {
                var c = zh[i];
                if (c >= '\u4e00' && c <= '\u9fff') {
                    chars.push(c);
                }
            }
        }

        var html = '<div class="pf-vocab-example">';
        html += '<div class="pf-vocab-example-label">📝 Câu ví dụ — Bấm vào từ để xem nghĩa</div>';
        html += '<div class="pf-vocab-example-zh">' + zh;
        if (zhJs && typeof speakText === 'function') {
            html += ' <button class="audio-btn-mini" '
                  + 'onclick="speakText(\'' + zhJs + '\', this, event)" '
                  + 'title="Nghe câu">'
                  + '<i class="fas fa-volume-up"></i>'
                  + '</button>';
        }
        html += '</div>';
        if (pinyin) {
            html += '<div class="pf-vocab-example-pinyin">'
                  + highlightPinyin(pinyin)
                  + '</div>';
        }
        if (vi) {
            html += '<div class="pf-vocab-example-vi">' + vi + '</div>';
        }

        if (chars.length > 0) {
            html += '<div class="pf-chars-label">👇 Bấm vào từ để xem nghĩa:</div>';
            html += '<div class="pf-chars-wrap">';
            chars.forEach(function(c, idx) {
                var cEsc = _esc(c);
                var cJs = String(c)
                    .replace(/\\/g, '\\\\')
                    .replace(/'/g, "\\'")
                    .replace(/"/g, '\\"');
                html += '<button class="pf-char-btn" '
                     + 'data-char="' + cEsc + '" '
                     + 'data-charjs="' + cJs + '" '
                     + 'onclick="vocabShowCharInfo(this, event)">'
                     + cEsc + '</button>';
            });
            html += '</div>';
            html += '<div class="pf-char-info" id="pfCharInfo" style="display:none"></div>';
        }

        html += '</div>';
        return html;
    }

    window._speakWord = function(text) {
        if (!text) return;
        if (!('speechSynthesis' in window)) return;
        try { speechSynthesis.cancel(); } catch(e) {}
        var u = new SpeechSynthesisUtterance(text);
        u.lang = 'zh-CN';
        if (typeof applyVoiceSettings === 'function') {
            try { applyVoiceSettings(u); } catch(e) {}
        }
        setTimeout(function() {
            try { speechSynthesis.speak(u); } catch(e) {}
        }, 30);
    };

    window._speakExampleKaraoke = function(exampleText) {
        if (!exampleText) return;
        if (!('speechSynthesis' in window)) return;
        try { speechSynthesis.cancel(); } catch(e) {}

        document.querySelectorAll('.pf-char-btn.reading').forEach(function(b) {
            b.classList.remove('reading');
        });
        var initFloat = document.getElementById('pfKaraokePinyin');
        if (initFloat) initFloat.remove();

        var examplePinyin = window.__currentVocabExamplePinyin || '';
        var pinyinWords = examplePinyin.trim().split(/\s+/).filter(function(w) { return w; });

        var chars = [];
        for (var i = 0; i < exampleText.length; i++) {
            var c = exampleText[i];
            if (c >= '\u4e00' && c <= '\u9fff') {
                chars.push(c);
            }
        }

        if (chars.length === 0) {
            var u0 = new SpeechSynthesisUtterance(exampleText);
            u0.lang = 'zh-CN';
            u0.rate = 0.75;
            if (typeof applyVoiceSettings === 'function') {
                try { applyVoiceSettings(u0); u0.rate = 0.75; } catch(e) {}
            }
            setTimeout(function() {
                try { speechSynthesis.speak(u0); } catch(e) {}
            }, 30);
            return;
        }

        var allBtns = document.querySelectorAll('.pf-char-btn');
        var btnChars = [];
        allBtns.forEach(function(b) {
            btnChars.push(b.dataset.char || '');
        });

        var list = [];
        if (window.FIXPY_DATASETS && window.FIXPY_DATASETS[VOCAB_ID]) {
            list = window.FIXPY_DATASETS[VOCAB_ID].data || [];
        }
        var pinyinMap = {};
        for (var m = 0; m < list.length; m++) {
            if (list[m].zh && list[m].pinyin && list[m].zh.length === 1) {
                pinyinMap[list[m].zh] = list[m].pinyin;
            }
        }

        var charIdx = 0;

        function clearFloatAndHighlight() {
            var f = document.getElementById('pfKaraokePinyin');
            if (f && f.parentNode) f.parentNode.removeChild(f);
            document.querySelectorAll('.pf-char-btn.reading').forEach(function(b) {
                b.classList.remove('reading');
            });
        }

        function speakNext() {
            if (charIdx >= chars.length) {
                setTimeout(function() {
                    clearFloatAndHighlight();
                }, 800);
                return;
            }

            var ch = chars[charIdx];

            clearFloatAndHighlight();

            var matchedBtn = null;
            for (var k = 0; k < btnChars.length; k++) {
                var btnChar = btnChars[k];
                if (btnChar === ch || btnChar.indexOf(ch) !== -1) {
                    matchedBtn = allBtns[k];
                    matchedBtn.classList.add('reading');
                    break;
                }
            }

            if (matchedBtn) {
                var btnCharM = matchedBtn.dataset.char || '';
                var pinyinText = '';

                for (var p = 0; p < list.length; p++) {
                    if (list[p].zh === btnCharM) {
                        pinyinText = list[p].pinyin || '';
                        break;
                    }
                }
                if (!pinyinText) pinyinText = pinyinMap[ch] || '';
                if (!pinyinText && pinyinWords.length === chars.length) {
                    pinyinText = pinyinWords[charIdx] || '';
                }
                if (!pinyinText && pinyinWords.length > 0) {
                    var ratio = charIdx / chars.length;
                    var pIdx = Math.floor(ratio * pinyinWords.length);
                    pinyinText = pinyinWords[Math.min(pIdx, pinyinWords.length - 1)] || '';
                }

                if (pinyinText) {
                    var rect = matchedBtn.getBoundingClientRect();
                    var floatEl = document.createElement('div');
                    floatEl.id = 'pfKaraokePinyin';
                    floatEl.className = 'pf-karaoke-pinyin';
                    floatEl.textContent = pinyinText;
                    floatEl.style.position = 'fixed';
                    floatEl.style.left = (rect.left + rect.width / 2) + 'px';
                    floatEl.style.top = (rect.top - 8) + 'px';
                    floatEl.style.transform = 'translate(-50%, -100%)';
                    floatEl.style.zIndex = '9999';
                    document.body.appendChild(floatEl);
                }
            }

            var u = new SpeechSynthesisUtterance(ch);
            u.lang = 'zh-CN';
            if (typeof applyVoiceSettings === 'function') {
                try { applyVoiceSettings(u); } catch(e) {}
            }
            u.rate = 0.75;
            if (!u.pitch || u.pitch < 0.8) u.pitch = 1.0;

            u.onend = function() {
                charIdx++;
                setTimeout(speakNext, 200);
            };
            u.onerror = function() {
                charIdx++;
                setTimeout(speakNext, 200);
            };

            setTimeout(function() {
                try {
                    speechSynthesis.speak(u);
                } catch(e) {
                    charIdx++;
                    setTimeout(speakNext, 200);
                }
            }, 30);
        }

        speakNext();
    };

    window.vocabShowCharInfo = function(btn, evt) {
        if (evt) { evt.stopPropagation(); evt.preventDefault(); }
        if (!btn) return;

        var char = btn.dataset.char || '';
        if (!char) return;

        var wasActive = btn.classList.contains('active');

        document.querySelectorAll('.pf-char-btn').forEach(function(b) {
            b.classList.remove('active', 'reading', 'zooming');
        });
        document.querySelectorAll('.similar-char-item').forEach(function(b) {
            b.classList.remove('active');
        });
        if ('speechSynthesis' in window) {
            try { speechSynthesis.cancel(); } catch(e) {}
        }

        var infoEl = document.getElementById('pfCharInfo');
        if (!infoEl) return;

        if (wasActive) {
            infoEl.style.display = 'none';
            return;
        }

        btn.classList.add('active');

        btn.classList.add('zooming');
        setTimeout(function() {
            btn.classList.remove('zooming');
        }, 600);

        _speakWord(char);

        var found = null;
        var list = [];
        if (window.FIXPY_DATASETS && window.FIXPY_DATASETS[VOCAB_ID]) {
            list = window.FIXPY_DATASETS[VOCAB_ID].data || [];
        }

        for (var i = 0; i < list.length; i++) {
            if (list[i].zh === char) {
                found = list[i];
                break;
            }
        }
        if (!found) {
            for (var j = 0; j < list.length; j++) {
                if (list[j].zh && list[j].zh.indexOf(char) !== -1) {
                    found = list[j];
                    break;
                }
            }
        }

        var html = '';

        html += '<div class="pf-char-info-header">';
        html += '<span class="pf-char-info-zh">' + _esc(char) + '</span>';

        if (found && found.pinyin) {
            html += '<span class="pf-char-info-pinyin">' + _esc(found.pinyin) + '</span>';
        }

        var exampleText = window.__currentVocabExampleZh || '';
        if (exampleText) {
            var exJs = String(exampleText)
                .replace(/\\/g, '\\\\')
                .replace(/'/g, "\\'")
                .replace(/"/g, '\\"');
            html += '<button class="pf-char-info-audio" '
                  + 'onclick="event.stopPropagation(); _speakExampleKaraoke(\'' + exJs + '\');" '
                  + 'title="Đọc cả câu ví dụ + karaoke">'
                  + '<i class="fas fa-volume-up"></i>'
                  + '</button>';
        }
        html += '</div>';

        if (found) {
            if (found.vi) {
                html += '<div class="pf-char-info-line">';
                html += '<span class="pf-char-info-label">📖 Nghĩa:</span>';
                html += '<span>' + _esc(found.vi) + '</span>';
                html += '</div>';
            }

            if (found.subject) {
                html += '<div class="pf-char-info-line">';
                html += '<span class="pf-char-info-label">🏷️ Loại từ:</span>';
                html += '<span>' + _esc(found.subject) + '</span>';
                html += '</div>';
            }

            if (found.radical) {
                var rad = found.radical;
                html += '<div class="pf-char-info-line">';
                html += '<span class="pf-char-info-label">🖌️ Bộ thủ:</span>';
                html += '<span>' + _esc(rad.zh || '');
                if (rad.pinyin) html += ' (' + _esc(rad.pinyin) + ')';
                if (rad.strokes) html += ' — ' + _esc(rad.strokes) + ' nét';
                if (rad.meaning) html += ' — ' + _esc(rad.meaning);
                html += '</span>';
                html += '</div>';
            }

            if (found.mnemonic) {
                html += '<div class="pf-char-info-line">';
                html += '<span class="pf-char-info-label">💡 Mẹo nhớ:</span>';
                html += '<span class="pf-char-info-mnemonic">'
                      + _esc(found.mnemonic).replace(/\n/g, '<br>')
                      + '</span>';
                html += '</div>';
            }
        } else {
            html += '<div class="pf-char-info-empty">';
            html += 'Không tìm thấy thông tin cho chữ "' + _esc(char) + '"';
            html += '</div>';
        }

        infoEl.innerHTML = html;
        infoEl.style.display = 'block';

        setTimeout(function() {
            infoEl.scrollIntoView({ behavior: 'smooth', block: 'nearest' });
        }, 100);
    };

    function setupPfViWatcher() {
        var pfViEl = document.getElementById('pfVi');
        if (!pfViEl) return;
        if (pfViEl.__vocabWatched) return;
        pfViEl.__vocabWatched = true;

        var observer = new MutationObserver(function() {
            if (!_isVocabMode()) return;
            if (!canAccessVocab()) return;

            var stt = null;
            try { stt = window.pfCurrentStt; } catch(e) {}
            if (!stt) return;

            var r = _findRecord(stt);
            if (!r || !r.vi_du_vi) return;

            var currentText = pfViEl.textContent.trim();
            if (currentText !== r.vi_du_vi) {
                pfViEl.textContent = r.vi_du_vi;
            }
        });

        observer.observe(pfViEl, {
            childList: true,
            characterData: true,
            subtree: true
        });
    }

    var _lastTier = null;
    function watchTier() {
        var key = 'always';
        if (key !== _lastTier) {
            _lastTier = key;
            updateTabLockState();
            bindTabIfNeeded();
        }
    }

    function init() {
        bindTabIfNeeded();
        updateTabLockState();
        setupObserver();
        setInterval(watchTier, 1000);
        hookPfBuildQuickNav();

        var tries = 0;
        var t = setInterval(function() {
            tries++;
            if (hookPracticeFull() || tries > 20) clearInterval(t);
        }, 250);
        hookPracticeFull();

        setTimeout(setupPfViWatcher, 1000);
        setInterval(setupPfViWatcher, 2000);
        setInterval(watchDatasetNav, 500);

        var _lastWatchedStt = null;
        var _lastWatchedDs = null;

        function watchSttChange() {
            if (!_isVocabMode()) {
                var blockR = document.querySelector('.pf-vocab-example');
                if (blockR) blockR.remove();

                var infoR = document.getElementById('pfCharInfo');
                if (infoR) {
                    infoR.style.display = 'none';
                    infoR.innerHTML = '';
                }

                _lastWatchedStt = null;
                _lastWatchedDs = null;
                return;
            }

            var curStt = (typeof window.pfCurrentStt !== 'undefined' && window.pfCurrentStt)
                         ? String(window.pfCurrentStt)
                         : null;

            var curDs = (typeof CURRENT_DATASET !== 'undefined')
                        ? CURRENT_DATASET
                        : 'tonghop';

            if (curStt === _lastWatchedStt && curDs === _lastWatchedDs) return;

            _lastWatchedStt = curStt;
            _lastWatchedDs = curDs;

            if (!curStt) return;

            var oldExample = document.querySelector('.pf-vocab-example');
            if (oldExample) oldExample.remove();

            var oldInfo = document.getElementById('pfCharInfo');
            if (oldInfo) {
                oldInfo.style.display = 'none';
                oldInfo.innerHTML = '';
            }

            var hintBtn = document.getElementById('pfHintBtn');
            if (hintBtn) {
                hintBtn.classList.remove('active');
                hintBtn.innerHTML = '<i class="fas fa-lightbulb"></i> Gợi ý';
            }
        }

        setInterval(watchSttChange, 250);
    }

    if (document.readyState === 'loading') {
        document.addEventListener('DOMContentLoaded', function() { setTimeout(init, 600); });
    } else {
        setTimeout(init, 600);
    }

    window.vocabUpdateLockState = updateTabLockState;

})();
"""
    js = js.replace("__SIMILAR_JSON__", similar_json)
    return js.replace("__VOCAB_ID__", vocab_id)
