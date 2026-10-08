from openpyxl import load_workbook
import re
import os
import shutil
from datetime import datetime

DATA_DIR = 'data'
FILE_NAME = 'tu_vung_hsk.xlsx'
BACKUP_PREFIX = 'tu_vung_hsk_old'
COLUMN = 'B'
START_ROW = 3

def normalize_text(text: str) -> str:
    """Chuẩn hóa cột B:
       - Dấu phân cách (kể cả fullwidth) -> |
       - Từ 1 đứng riêng, các từ sau viết liền
       - Kết quả: 'A | BC' hoặc 'A | B' hoặc 'A'
    """
    if not text or not isinstance(text, str):
        return text

    text = text.strip()

    # 1. Dấu phân cách: ASCII + fullwidth
    cleaned = re.sub(
        r'[\(\)\[\]\{\};,/\-_'
        r'（）［］｛｝；，、／＼－＿—]+',
        '|', text
    )

    # 2. Khoảng trắng quanh | -> |
    cleaned = re.sub(r'\s*\|\s*', '|', cleaned)

    # 3. Gộp nhiều | thành 1
    cleaned = re.sub(r'\|+', '|', cleaned)

    # 4. Bỏ | đầu/cuối
    cleaned = cleaned.strip('|')

    # 5. Nếu không có | mà có khoảng trắng giữa chữ Hán -> coi là phân cách
    if '|' not in cleaned and re.search(r'[\u4e00-\u9fff]\s+[\u4e00-\u9fff]', cleaned):
        cleaned = re.sub(r'\s+', '|', cleaned)

    # 6. Tách, lọc rỗng
    parts = [p.strip() for p in cleaned.split('|') if p.strip()]

    if not parts:
        return text
    if len(parts) == 1:
        return parts[0]

    # 7. Từ đầu riêng, các từ sau viết liền
    first = parts[0]
    rest = ''.join(parts[1:])   # ⚠️ nối liền không dấu cách

    return f'{first} | {rest}'


def verify_only_column_b_changed(src, dst, column='B'):
    print('\n🔍 Kiểm tra an toàn...')
    wb_src = load_workbook(src)
    wb_dst = load_workbook(dst)

    if wb_src.sheetnames != wb_dst.sheetnames:
        print('❌ Danh sách sheet khác nhau!')
        return False

    diffs_outside_b = 0
    diffs_inside_b = 0

    for sheet in wb_src.sheetnames:
        ws1, ws2 = wb_src[sheet], wb_dst[sheet]
        for row in ws1.iter_rows():
            for c1 in row:
                c2 = ws2[c1.coordinate]
                if c1.value != c2.value:
                    if c1.column_letter == column:
                        diffs_inside_b += 1
                    else:
                        diffs_outside_b += 1
                        print(f'   ⚠️ Khác ngoài cột B: {sheet}!{c1.coordinate}: "{c1.value}" → "{c2.value}"')

    print(f'\n📊 Kết quả kiểm tra:')
    print(f'   - Số ô cột {column} đã sửa: {diffs_inside_b}')
    print(f'   - Số ô NGOÀI cột {column} bị thay đổi: {diffs_outside_b}')

    if diffs_outside_b == 0:
        print(f'   ✅ Xác nhận: chỉ cột {column} thay đổi, các cột khác nguyên vẹn.')
        return True
    else:
        print(f'   ❌ CẢNH BÁO: có {diffs_outside_b} ô ngoài cột {column} bị thay đổi!')
        return False


def main():
    src_file = os.path.join(DATA_DIR, FILE_NAME)

    if not os.path.isfile(src_file):
        print(f'❌ Không tìm thấy file: {src_file}')
        return

    tmp_file = os.path.join(DATA_DIR, '_tmp_normalized.xlsx')
    shutil.copyfile(src_file, tmp_file)
    print(f'📋 Đã copy file gốc → {tmp_file}')

    wb = load_workbook(tmp_file)
    total_edit = 0

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f'  📄 Sheet: {sheet_name}')
        for row in range(START_ROW, ws.max_row + 1):
            cell = ws[f'{COLUMN}{row}']
            old_value = cell.value
            if old_value is None:
                continue
            new_value = normalize_text(str(old_value))
            if new_value != str(old_value).strip():
                cell.value = new_value
                total_edit += 1
                print(f'     ✔ {COLUMN}{row}: "{old_value}" → "{new_value}"')

    wb.save(tmp_file)
    print(f'\n✅ Đã sửa {total_edit} ô trong file tạm.')

    if not verify_only_column_b_changed(src_file, tmp_file, COLUMN):
        os.remove(tmp_file)
        raise SystemExit('❌ Kiểm tra thất bại, hủy thao tác!')

    backup_file = os.path.join(DATA_DIR, f'{BACKUP_PREFIX}.xlsx')
    if os.path.exists(backup_file):
        ts = datetime.now().strftime('%Y%m%d_%H%M%S')
        backup_file = os.path.join(DATA_DIR, f'{BACKUP_PREFIX}_{ts}.xlsx')

    shutil.move(src_file, backup_file)
    print(f'\n📦 Đã backup file gốc → {backup_file}')

    shutil.move(tmp_file, src_file)
    print(f'🔄 File chuẩn hóa đã thay thế → {src_file}')

    print('\n🎉 Hoàn tất!')


if __name__ == '__main__':
    main()
